#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
# Copyright (c) 2026 FR. All rights reserved.
# 本文件是 FXseek 的一部分：非商业用途免费，商业用途需取得授权（见仓库根目录 COMMERCIAL.md）。

"""
本地语义检索核心引擎（图片 / 视频 / 文档 统一向量索引 + 自然语言搜索）

★ 依赖 wemm_app/venv 独立环境（mlx + mlx_vlm + PIL + sqlite3），不碰 oMLX。
★ 零额外第三方依赖：sqlite3 存向量，ffmpeg 抽视频帧，textutil 转文档。

用法：
    # 建索引（可多次增量）
    venv/cpython-3.11/bin/python3.11 indexer.py index ~/Pictures ~/Movies ~/Documents
    # 搜索
    venv/cpython-3.11/bin/python3.11 indexer.py search "一只在操场上跑的狗"
    # 查看索引统计
    venv/cpython-3.11/bin/python3.11 indexer.py stats
"""
import argparse
import hashlib
import json
import os
import sqlite3
import struct
import subprocess
import sys
import tempfile
import time
from pathlib import Path

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import heif_support  # noqa: F401  注册 HEIC/HEIF 解码器（iPhone 照片），进程级生效
import paths as P
DB_PATH = P.DB_PATH

# ★ 允许解码「尾部略有截断」的图片。
#   老唱片封面这类素材常有几十字节缺失（Pillow 报 image file is truncated），
#   但画面本身几乎完好。直接当坏图跳过有两个坏处：
#     ① 明明能看的图搜不到；
#     ② 跳过后 items 里没有它的记录 → 每次刷新都被重新当成「新文件」重试，
#        报告里就会出现「新索引 1 个文件 / 0 个向量」这种看着像 bug 的假消息。
try:
    from PIL import ImageFile as _PILImageFile
    _PILImageFile.LOAD_TRUNCATED_IMAGES = True
except Exception:
    pass

IMAGE_EXT = {".jpg", ".jpeg", ".png", ".gif", ".bmp", ".webp", ".tiff", ".heic"}
VIDEO_EXT = {".mp4", ".mov", ".avi", ".mkv", ".m4v", ".webm", ".flv", ".wmv"}
AUDIO_EXT = {".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".opus", ".aiff"}
DOC_EXT = {".txt", ".md", ".pdf", ".doc", ".docx", ".rtf", ".wps", ".csv", ".json", ".py",
           ".js", ".ts", ".html", ".htm", ".xml", ".yaml", ".yml", ".log",
           # 表格（openpyxl 读单元格文本；.et 见 _read_text 的双路径处理）
           ".xlsx", ".xlsm", ".et",
           # 幻灯片（python-pptx 读文本框；.dps 见 _read_text 的降级处理）
           ".pptx", ".dps"}

# 系统生成的垃圾/元数据文件，不是真素材。
# ★ 为什么必须显式过滤：macOS 在非 HFS 卷和网络卷上会给每个文件配一个 AppleDouble
#   伴生文件 "._原名"，它**保留原扩展名**（._03 Hurt You.m4a），于是能通过扩展名白名单
#   被当音频索引，编码时 ffmpeg 直接报 "moov atom not found"。
JUNK_NAMES = {".DS_Store", ".localized", "Thumbs.db", "desktop.ini", "ehthumbs.db"}


def is_junk_file(name: str) -> bool:
    """该文件名是否属于「不该入库」的垃圾/元数据文件（只看文件名，不看路径）。"""
    if name.startswith("._"):   # AppleDouble 伴生文件
        return True
    if name.startswith("."):    # .DS_Store 等一切点文件
        return True
    return name in JUNK_NAMES


def purge_stale(cur, con, roots):
    """清理索引中已失效的记录，返回删除条数。

    要删两类：
      ① 源目录里已经不存在的文件；
      ② 垃圾/元数据文件 —— 旧版本会把 "._03 Hurt You.m4a" 这类 AppleDouble
         当音频索引进库（它保留了原扩展名，能通过白名单）。
    只清理本轮 roots 覆盖范围内的记录，避免误删其他索引源的数据。
    """
    def norm(x):
        try:
            return os.path.realpath(os.path.expanduser(x)).rstrip(os.sep)
        except Exception:
            return None

    def under(p, rs):
        try:
            rp = os.path.realpath(p)
        except Exception:
            rp = p
        return any(rp == r or rp.startswith(r + os.sep) for r in rs)

    mine = {r for r in (norm(x) for x in roots) if r}
    # ★ 源目录整体不可达时（外接盘没插、网络盘没挂）**不能**按「文件不存在」删记录，
    #   否则拔一次盘、再索引一次，整个源的索引就被清空了。
    #   垃圾文件不受此限：只看文件名就能判定，删掉永远是对的。
    reachable = {r for r in mine if os.path.exists(r)}

    n_removed = 0
    for (p,) in cur.execute("SELECT DISTINCT path FROM items").fetchall():
        if not under(p, mine):
            continue
        if is_junk_file(os.path.basename(p)):
            cur.execute("DELETE FROM items WHERE path=?", (p,))
            n_removed += 1
        elif under(p, reachable) and not os.path.exists(p):
            cur.execute("DELETE FROM items WHERE path=?", (p,))
            n_removed += 1
    if n_removed:
        con.commit()
    return n_removed


# ffmpeg / ffprobe：优先用 app 内置（bin/ 目录，随应用分发），其次环境变量，最后系统
def _find_bin(name: str) -> str:
    """按优先级查找可执行文件：内置 > 环境变量 > 系统 PATH。"""
    import shutil
    # 1) app 内置（打包分发用，目标机器无需安装 ffmpeg）
    for cand in (os.path.join(HERE, "bin", name),
                 os.path.join(HERE, "bin", name + ".exe")):
        if os.path.exists(cand) and os.access(cand, os.X_OK):
            return cand
    # 2) 环境变量覆盖
    env = os.environ.get(name.upper())
    if env and os.path.exists(env):
        return env
    # 3) 系统 PATH / 常见位置
    for cand in (shutil.which(name) or "",
                 f"/opt/homebrew/bin/{name}", f"/usr/local/bin/{name}",
                 f"/usr/bin/{name}"):
        if cand and os.path.exists(cand):
            return cand
    return name

FFMPEG = _find_bin("ffmpeg")
FFPROBE = _find_bin("ffprobe")


# ===================== 媒体元数据缓存 =====================
# 为什么需要：probe_media_meta 对每个音视频要起一次 ffprobe 子进程（实测 50~110ms）、
# 图片要开一次 PIL（~40ms）。272 个文件串行下来就是 ~19 秒 —— 用户看到的
# 「加载素材…」几乎全花在这儿（实测 POST /v1/browse 18.8s，接口本身只做了一次查询）。
# 元数据只取决于文件内容，所以用 (路径, mtime, 大小) 当 key：
# 文件一改 key 就变，天然不会读到过期值；没动过的文件永远不用再 probe。
_PROBE_MEM = {}          # key -> meta（本进程用过）
_PROBE_DISK = {}         # key -> [写入时间, meta]（跨次启动复用）
_PROBE_LOADED = False
_PROBE_DIRTY = 0
_PROBE_LAST_SAVE = 0.0
_PROBE_MAX = 6000        # 磁盘缓存上限条数，超了按写入时间淘汰
import threading as _threading
_PROBE_LOCK = _threading.Lock()
try:
    _PROBE_FILE = os.path.join(P.DATA_DIR, "media_meta_cache.json")
except Exception:
    _PROBE_FILE = os.path.join(HERE, "media_meta_cache.json")


def _probe_key(path):
    """(路径, mtime, 大小) —— 拿不到 stat 就不缓存（文件可能已不在）。

    前缀 v2：探测内容变过（新增容器标签 tags）时必须让旧缓存失效，
    否则老缓存会一直返回「没有 tags」的旧结果，界面永远等不到歌手/专辑。
    以后只要改 _probe_media_meta_raw 提取的字段，就把这个版本号 +1。
    """
    try:
        st = os.stat(path)
    except OSError:
        return None
    return "v2|%s|%d|%d" % (path, int(st.st_mtime), st.st_size)


def _probe_load_disk():
    global _PROBE_LOADED
    if _PROBE_LOADED:
        return
    _PROBE_LOADED = True
    try:
        with open(_PROBE_FILE, encoding="utf-8") as f:
            d = json.load(f)
        if isinstance(d, dict):
            _PROBE_DISK.update(d)
    except Exception:
        pass


def _probe_save_disk(force=False):
    """写盘是「攒一批再写」，避免每 probe 一个文件就落一次盘。"""
    global _PROBE_DIRTY, _PROBE_LAST_SAVE, _PROBE_DISK, _PROBE_MEM
    if not _PROBE_DIRTY:
        return
    now = time.time()
    if not force and (now - _PROBE_LAST_SAVE) < 15:
        return
    with _PROBE_LOCK:
        _PROBE_LAST_SAVE = now
        _PROBE_DIRTY = 0
        merged = dict(_PROBE_DISK)
        for k, v in _PROBE_MEM.items():
            merged[k] = [now, v]
        if len(merged) > _PROBE_MAX:
            keep = sorted(merged.items(), key=lambda kv: kv[1][0] if isinstance(kv[1], list) else 0)
            merged = dict(keep[-_PROBE_MAX:])
        try:
            tmp = _PROBE_FILE + ".tmp"
            with open(tmp, "w", encoding="utf-8") as f:
                json.dump(merged, f, ensure_ascii=False)
            os.replace(tmp, _PROBE_FILE)
        except Exception:
            pass


def _probe_uncached(p):
    try:
        return probe_media_meta(p)
    except Exception:
        return {}


def warm_meta_cache(paths, workers=8):
    """并行预热：只对「缓存里没有」的文件真的跑 probe，返回预热条数。

    串行 19s → 8 线程约 2~3s（probe 基本都是等子进程，GIL 不是瓶颈），
    且跑过一次之后同样的文件直接命中缓存（毫秒级）。
    """
    _probe_load_disk()
    todo = []
    for p in paths:
        k = _probe_key(p)
        if k and k not in _PROBE_MEM and k not in _PROBE_DISK:
            todo.append(p)
    if not todo:
        return 0
    workers = max(1, min(int(workers or 8), len(todo)))
    if workers == 1:
        for p in todo:
            _probe_uncached(p)
    else:
        try:
            from concurrent.futures import ThreadPoolExecutor
            with ThreadPoolExecutor(max_workers=workers) as ex:
                list(ex.map(_probe_uncached, todo))
        except Exception:
            for p in todo:
                _probe_uncached(p)
    _probe_save_disk(force=True)
    return len(todo)


def probe_media_meta(path: str) -> dict:
    """带缓存的探测入口（缓存逻辑见上方 _PROBE_MEM 注释）。

    ★ 一律返回**副本**：缓存里那个 dict 是全局共享的，调用方（比如
    app.file_detail 把索引里的 asr_text/ai_tags 并进去）一旦就地改它，
    改动会永远留在缓存里 —— 之后哪怕库里的字段已经删掉，探测仍然吐旧值，
    而且下一次 _probe_save_disk 会把这个脏数据写进磁盘缓存。
    """
    meta = _probe_media_meta_cached(path)
    return dict(meta) if isinstance(meta, dict) else meta


def _probe_media_meta_cached(path: str) -> dict:
    global _PROBE_DIRTY
    k = _probe_key(path)
    if k:
        hit = _PROBE_MEM.get(k)
        if hit is not None:
            return hit
    _probe_load_disk()
    if k:
        disk = _PROBE_DISK.get(k)
        if disk is not None:
            meta = disk[1] if isinstance(disk, list) and len(disk) == 2 else disk
            _PROBE_MEM[k] = meta
            return meta
    meta = _probe_media_meta_raw(path)
    if k:
        _PROBE_MEM[k] = meta
        _PROBE_DIRTY += 1
        _probe_save_disk()
    return meta


def _fix_tag_mojibake(v: str) -> str:
    """修 ID3 里最常见的乱码：中文标签按 GBK 存、播放器却当 latin-1 解出来。

    实测 `Bruno Mars - Lazy Song.mp3` 的 artist 是 `|ÆßÉ«³±ÒôÉç|in7se.com|Æâse|`，
    其实是「七色潮音社」那类中文被错解。这类字符串**不能原样进向量**（等于往库里
    灌噪声），也**不能原样显示**（卡片上一片乱码）。所以先试着按 GBK/BIG5 重解一次，
    解不出中文就返回空串让调用方丢弃 —— 宁可不显示，也不要乱码。
    """
    if not v:
        return ""
    # 没有 latin-1 高位字符 = 不是这种乱码，原样返回
    if not any("\u00c0" <= ch <= "\u00ff" for ch in v):
        return v
    for enc in ("gbk", "gb18030", "big5", "shift_jis"):
        try:
            fixed = v.encode("latin-1").decode(enc)
        except Exception:
            continue
        cjk = sum(1 for ch in fixed if "\u4e00" <= ch <= "\u9fff")
        if cjk >= 2 and not any("\u00c0" <= ch <= "\u00ff" for ch in fixed):
            return fixed
    return ""          # 修不好 → 丢掉，别把乱码写进向量和界面


def _probe_media_meta_raw(path: str) -> dict:
    """探测媒体元数据：视频时长、图片尺寸、文档字数、音频时长。（无缓存版）

    用于 UI 卡片角标显示（如视频右下角 00:15）。
    """
    ext = os.path.splitext(path)[1].lower()
    meta = {}
    try:
        if ext in VIDEO_EXT or ext in AUDIO_EXT:
            # ★ 只跑一次 ffprobe：-show_format 里本来就带 duration，
            #   以前先跑 _probe_duration 再跑一次完整信息 = 每个文件两次子进程（白花一倍时间）。
            info = {}
            try:
                out = subprocess.run(
                    [FFPROBE, "-v", "error", "-show_format", "-show_streams",
                     "-of", "json", path],
                    capture_output=True, text=True, timeout=25)
                info = json.loads(out.stdout or "{}")
            except Exception:
                info = {}
            fmt = info.get("format", {}) or {}
            try:
                dur = float(fmt.get("duration") or 0)
            except Exception:
                dur = 0.0
            if dur <= 0:
                dur = _probe_duration(path)      # 兜底：个别容器 format 里没有 duration
            if dur > 0:
                meta["duration"] = round(dur, 1)
                meta["duration_text"] = _fmt_duration(dur)
            # ★ 音乐文件的容器标签（ID3 / Vorbis / MP4 atoms）。
            #   ffprobe -show_format 本来就把 tags 一起返回来 —— 以前只取 duration /
            #   bit_rate，tags 直接白扔了。而「歌手/专辑/曲名」是文件名之外**唯一**
            #   的语义来源：别人的库里文件名常常是 01.mp3、Track 03.flac，
            #   没有标签就等于这首歌搜不到歌手、搜不到专辑。
            raw_tags = fmt.get("tags") or {}
            if raw_tags:
                tl = {str(k).lower(): v for k, v in raw_tags.items()}

                def _tag(*names, limit=120):
                    for nm in names:
                        v = tl.get(nm)
                        if v is None:
                            continue
                        v = str(v).strip()
                        if not v or v.lower() in ("unknown", "unknown artist",
                                                  "unknown album", "various artists",
                                                  "n/a", "-"):
                            continue
                        # 盗版站常把广告塞进 album/genre/composer（实测
                        # "www.alexak.ro" 同时占了四个字段）。带网址的一律丢掉，
                        # 否则卡片上会显示「Dj Project feat. Giulia · www.alexak.ro」。
                        low = v.lower()
                        if "http" in low or "www." in low or ".ru/" in low or ".ro/" in low:
                            continue
                        v = _fix_tag_mojibake(v)
                        if not v:            # 乱码修不好 → 丢弃
                            continue
                        # 重解之后可能又露出网址（乱码把 .com 藏起来了），再查一次
                        low = v.lower()
                        if "http" in low or "www." in low or ".com" in low or ".net" in low:
                            continue
                        return v[:limit]
                    return ""

                tg = {
                    "title":        _tag("title"),
                    "artist":       _tag("artist"),
                    "album":        _tag("album"),
                    "album_artist": _tag("album_artist", "albumartist"),
                    "genre":        _tag("genre", limit=40),
                    "track":        _tag("track", limit=12),
                    "composer":     _tag("composer", limit=80),
                    "date":         _tag("date", "year", limit=24),
                }
                tg = {k: v for k, v in tg.items() if v}
                # 年份单独抽出来（只用于显示）：标签里常有脏值（实测见过 5583），
                # 所以只认 1900–2099 的四位前缀。
                d0 = (tg.get("date") or "")[:4]
                if d0.isdigit() and 1900 <= int(d0) <= 2099:
                    tg["year"] = d0
                if tg:
                    meta["tags"] = tg
                # 不管有没有可用标签，都记一笔「已查过」：既是迁移标记
                # （老记录没这个键 → 重编一次），也避免每次刷新再跑一遍 ffprobe。
                meta["tags_checked"] = True
            try:
                meta["container"] = (fmt.get("format_long_name")
                                     or fmt.get("format_name") or "").split(",")[0]
                if fmt.get("bit_rate"):
                    try: meta["bitrate_kbps"] = round(int(fmt["bit_rate"]) / 1000)
                    except Exception: pass
                vstreams = [s for s in info.get("streams", [])
                            if s.get("codec_type") == "video"]
                if vstreams:
                    v = vstreams[0]
                    if v.get("width") and v.get("height"):
                        meta["width"], meta["height"] = int(v["width"]), int(v["height"])
                        meta["dim_text"] = f'{v["width"]}×{v["height"]}'
                    meta["vcodec"] = (v.get("codec_long_name") or v.get("codec_name") or "")
                    if v.get("pix_fmt"): meta["pix_fmt"] = v["pix_fmt"]
                    if v.get("profile"): meta["vprofile"] = v["profile"]
                    fr = v.get("avg_frame_rate") or v.get("r_frame_rate") or ""
                    if "/" in fr:
                        try:
                            num, den = fr.split("/")
                            if float(den) > 0:
                                meta["fps"] = round(float(num) / float(den), 2)
                        except Exception: pass
                astreams = [s for s in info.get("streams", [])
                            if s.get("codec_type") == "audio"]
                if astreams:
                    a = astreams[0]
                    meta["acodec"] = (a.get("codec_long_name") or a.get("codec_name") or "")
                    if a.get("sample_rate"): meta["sample_rate"] = int(a["sample_rate"])
                    if a.get("channels"): meta["channels"] = int(a["channels"])
            except Exception:
                pass
        elif ext in IMAGE_EXT:
            from PIL import Image
            with Image.open(path) as im:
                meta["width"], meta["height"] = im.size
                meta["dim_text"] = f"{im.size[0]}×{im.size[1]}"
                meta["img_format"] = im.format or ext.lstrip(".").upper()
                meta["img_mode"] = im.mode
                try:
                    dpi = im.info.get("dpi")
                    if dpi: meta["dpi"] = f"{int(dpi[0])}×{int(dpi[1])}"
                except Exception:
                    pass
                try:
                    if getattr(im, "n_frames", 1) > 1:
                        meta["frames"] = im.n_frames
                except Exception:
                    pass
        elif ext in DOC_EXT:
            try:
                txt = ""
                if _is_probably_text(path):
                    with open(path, "r", encoding="utf-8", errors="ignore") as f:
                        txt = f.read()
                lines = len(txt.splitlines()) if txt else 0
                chars = len(txt)
                meta["lines"] = lines
                meta["chars"] = chars
                meta["words"] = len(txt.split()) if txt else 0
                meta["dim_text"] = f"{chars} 字" if chars else ""
                with open(path, "rb") as f:
                    raw = f.read(8192)
                for enc in ("utf-8", "gbk", "big5", "latin-1"):
                    try:
                        raw.decode(enc); meta["encoding"] = enc; break
                    except Exception:
                        continue
            except Exception:
                pass
        # 通用：文件大小 + 修改时间
        try:
            st = os.stat(path)
            meta["size"] = st.st_size
            meta["size_text"] = _fmt_size(st.st_size)
            meta["mtime"] = st.st_mtime
        except OSError:
            pass
    except Exception:
        pass
    return meta


def _fmt_duration(sec: float) -> str:
    s = int(round(sec))
    h, rem = divmod(s, 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _fmt_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024.0
    return f"{n:.1f} GB"


def ffmpeg_status() -> dict:
    """ffmpeg 运行环境自检（判断是否用内置、能否正常执行）。"""
    import subprocess
    builtin = os.path.join(HERE, "bin", "ffmpeg")
    info = {
        "ffmpeg": FFMPEG,
        "ffprobe": FFPROBE,
        "builtin": os.path.exists(builtin),
        "using_builtin": os.path.realpath(FFMPEG) == os.path.realpath(builtin),
        "ok": False,
        "version": None,
    }
    try:
        out = subprocess.run([FFMPEG, "-version"], capture_output=True,
                             text=True, timeout=15)
        if out.returncode == 0 and out.stdout:
            info["ok"] = True
            info["version"] = out.stdout.splitlines()[0].strip()
    except Exception as e:
        info["error"] = str(e)
    return info

# 视频抽帧策略
# ---- 视频抽帧自适应策略 ----
# 按时长分段：短视频密抽（定位准），长视频疏抽（控资源）
# 每项：(时长上限秒, 抽帧间隔秒, 帧数上限)
VIDEO_FRAME_TIERS = [
    (10,   0.5,  24),    # ≤10s：每 0.5s 一帧
    (30,   1.0,  24),    # ≤30s：每 1s 一帧
    (120,  2.0,  32),    # ≤2min：每 2s 一帧
    (600,  5.0,  48),    # ≤10min：每 5s 一帧
    (float("inf"), 10.0, 60),   # >10min：每 10s 一帧
]
VIDEO_MAX_FRAMES = 32     # 全局兜底上限（可被设置覆盖）

# 文档分块
DOC_CHUNK_CHARS = 800
DOC_CHUNK_OVERLAP = 100
# 「空文档」判定：正文短于此长度的文档，语义向量等于一个空文档，对任何查询
# 都有 0.5~0.64 的相似度（实测同库里 7 个这样的文件恒定霸占所有查询的前 5 名）。
EMPTY_TEXT_LEN = 20
# 空文档的语义分惩罚系数（只罚语义分，不罚文件名命中的词法加成）。
EMPTY_DOC_PENALTY = 0.35

# --------------------------------------------------------------------------
# AI 标签向量（chunk_idx >= AI_TAG_BASE）
# --------------------------------------------------------------------------
# 图片在索引期由视觉模型生成 ai_tags，每个标签单独编码成一个向量存进库里，
# 让「文字描述画面」变成 text→text（否则纯图像向量对文字查询只有 0.38 左右，
# 和无关内容的背景分 0.42 重叠，任何阈值都分不开）。
#
# 但短词向量有个通病：两个任意短词之间的余弦就有 0.5~0.62（实测 32 个标签
# 的均值 0.44~0.62），直接当分数用会让「太空」「财务」这类查不到东西的词也
# 命中一堆图片。这里用「减去通用标签方向」的办法居中，居中后的点积实测：
#   真匹配（查询就是标签本身）  0.454 ~ 0.522
#   句子式查询（「有气球的画面」）0.286 ~ 0.478
#   无关查询（太空/财务/会议…） 0.000 ~ 0.097
# 中间的空档足够宽，配合 fastsearch.map_tag_score 映射成与文档可比的分数。
AI_TAG_BASE = 9000        # 标签向量起始 chunk_idx
AI_DESC_IDX = 9500        # AI 描述句向量（每个文件一条）
# 用来定义「通用标签方向」的词表：取一批最常见、信息量最低的标签词求平均。
TAG_CENTER_TERMS = ["人物", "自然", "户外", "场景", "物体", "内容", "图片", "照片",
                    "天空", "建筑", "食物", "动物", "植物", "室内", "室外", "颜色",
                    "装饰", "背景", "风景", "日常", "生活", "环境"]
_TAG_CENTER = None        # 懒加载缓存


def tag_center_vec(model, processor):
    """「通用标签方向」的平均向量（首次调用时编码并缓存）。"""
    global _TAG_CENTER
    if _TAG_CENTER is None:
        import embed as we
        import mlx.core as mx
        import numpy as _np
        acc = None
        for t in TAG_CENTER_TERMS:
            v = we.embed(model, processor, text=t, instruction=we.DOC_INSTRUCTION)
            mx.eval(v)
            a = _np.array(v.tolist(), dtype=_np.float32)
            acc = a if acc is None else acc + a
        _TAG_CENTER = acc / len(TAG_CENTER_TERMS)
    return _TAG_CENTER


def _center_text_vec(model, processor, text):
    """把一段文字编码成「去掉通用标签方向」后的单位向量。"""
    import embed as we
    import mlx.core as mx
    import numpy as _np
    v = we.embed(model, processor, text=text, instruction=we.DOC_INSTRUCTION)
    mx.eval(v)
    a = _np.array(v.tolist(), dtype=_np.float32) - tag_center_vec(model, processor)
    n = float(_np.linalg.norm(a))
    if n < 1e-6:
        return None
    return (a / n).tolist()


def ai_tag_status(db_path=None):
    """图片/视频里，哪些已经有 AI 标签向量、哪些还没有。

    返回 {total, tagged, pending: [{path, kind, name}]}。
    「已打标」的判据是**有 chunk_idx >= AI_TAG_BASE 的向量行**，不是看 meta —— 这样
    不会因为手改过 meta 而误判。pending 顺序跟浏览列表一致（按首次入库的 rowid）。
    """
    db = db_path or DB_PATH
    con = sqlite3.connect(db)
    try:
        tagged = {p for (p,) in con.execute(
            "SELECT DISTINCT path FROM items WHERE chunk_idx>=?", (AI_TAG_BASE,))}
        rows = con.execute(
            "SELECT path, kind FROM items"
            " WHERE kind IN ('image','video') AND chunk_idx<?"
            " GROUP BY path, kind ORDER BY MIN(id)", (AI_TAG_BASE,)).fetchall()
    except sqlite3.OperationalError:
        rows = []
        tagged = set()
    finally:
        con.close()
    pend = [(p, k) for (p, k) in rows if p not in tagged]
    return {"total": len(rows), "tagged": len(rows) - len(pend),
            "pending": [{"path": p, "kind": k, "name": os.path.basename(p)}
                        for (p, k) in pend]}


def _merge_meta(con, path, patch, clear=()):
    """把一段 meta 并进某个文件所有「非 AI 槽位」的行（AI 标签向量行不动）。

    clear：本次要**删掉**的键。dict.update 只能覆盖、不能删除，所以「重新转写」
    这次没识别到文字时必须显式清掉上一次的 asr_text —— 否则旧歌词会永远留着，
    用户重转一次反而还看得到老内容（本次要修的就是这个）。
    """
    for ci, mstr in con.execute(
            "SELECT chunk_idx, meta FROM items WHERE path=? AND chunk_idx<?",
            (path, AI_TAG_BASE)).fetchall():
        try:
            m = json.loads(mstr) if mstr else {}
        except Exception:
            m = {}
        m.update(patch)
        for k in clear:
            m.pop(k, None)
        con.execute("UPDATE items SET meta=? WHERE path=? AND chunk_idx=?",
                    (json.dumps(m, ensure_ascii=False), path, ci))


def tag_file(path, model, processor, ai_cfg, db_path=None):
    """给单个图片/视频补做 AI 描述与标签，**只碰 AI 那几条向量**。

    这是「详情页立即打标签」和「闲置自动打标」共用的轻量路径：不像
    `build_index(force=True)` 那样把整份文件重新编码（视频要重抽帧、图片要重编码），
    只做三件事 —— 生成描述、把 ai_meta 并进该文件已有行的 meta、重建 9000+ 的标签向量。
    所以视频打标从「几十秒编码 N 帧」降到「一次视觉模型调用 + M 次文本编码」。

    返回 {ok, tags, description, vectors, seconds} 或 {ok: False, message}。
    """
    import ai_desc as _ai
    from PIL import Image as _PILImage

    db = db_path or DB_PATH
    ext = os.path.splitext(path)[1].lower()
    if ext not in (IMAGE_EXT | VIDEO_EXT):
        return {"ok": False, "message": "只支持图片与视频"}
    if not ai_cfg or not (ai_cfg.get("base_url") and ai_cfg.get("model")):
        return {"ok": False, "message": "请先在设置里配置 AI 视觉模型与 API 地址"}
    if not os.path.exists(path):
        return {"ok": False, "message": "文件不存在"}

    con = sqlite3.connect(db)
    try:
        row = con.execute(
            "SELECT kind, mtime FROM items WHERE path=? AND chunk_idx=0",
            (path,)).fetchone()
        if not row:
            return {"ok": False, "message": "这个素材还没有索引，先刷新索引再打标"}
        kind, mtime = row[0], row[1]

        t0 = time.time()
        # 与索引期完全一致的送图方式：图片＝这一张；视频＝关键帧九宫格
        if ext in IMAGE_EXT:
            with _PILImage.open(path) as im:
                pic = im.convert("RGB")
            prompt = ai_cfg.get("prompt", "")
            kind = "image"
        else:
            pic = _video_contact_sheet(path)
            prompt = ai_cfg.get("prompt_video") or VIDEO_SHEET_PROMPT
            kind = "video"

        desc = _ai.describe_image(ai_cfg["base_url"], ai_cfg.get("api_key", ""),
                                  ai_cfg["model"], pic, prompt=prompt,
                                  max_tags=int(ai_cfg.get("max_tags", 12) or 12))
        tags = [t for t in (desc.get("tags") or []) if t]
        text = (desc.get("description") or "").strip()
        ai_meta = {"ai_description": text, "ai_tags": tags,
                   "ai_model": ai_cfg["model"]}

        # 1) 把 ai_meta 并进该文件所有「非 AI」行的 meta（详情页读的就是 chunk 0 那条）
        _merge_meta(con, path, ai_meta)

        # 2) 重建标签向量：先删旧的，避免新标签变少时留下搜得到的孤儿向量
        con.execute("DELETE FROM items WHERE path=? AND chunk_idx>=?",
                    (path, AI_TAG_BASE))
        n_vec = 0
        slots = [(AI_TAG_BASE + i, t, True) for i, t in enumerate(tags)]
        if text:
            slots.append((AI_DESC_IDX, text, False))
        for si, stext, is_tag in slots:
            try:
                cv = _center_text_vec(model, processor, stext)
            except Exception as e:
                print(f"  [AI标签编码失败] {path}: {e}")
                continue
            if not cv:
                continue
            tmeta = dict(ai_meta)
            tmeta["_ai_tag_slot" if is_tag else "_ai_desc_slot"] = True
            con.execute(
                "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                " VALUES (?,?,?,?,?,?,?,?)",
                (path, kind, si, json.dumps(tmeta, ensure_ascii=False), len(cv),
                 _to_blob(cv), mtime or time.time(), time.time()))
            n_vec += 1
        con.commit()
        return {"ok": True, "tags": tags, "description": text,
                "vectors": n_vec, "seconds": round(time.time() - t0, 1)}
    except _ai.AIError as e:
        return {"ok": False, "message": str(e)}
    except Exception as e:
        return {"ok": False, "message": f"{type(e).__name__}: {e}"}
    finally:
        con.close()


def _has_audio_stream(path: str) -> bool:
    """ffprobe 看一下有没有音频流。无声视频（测试卡、纯画面录像）不必送去转写 ——
    ASR 服务会回一个 HTTP 500「No audio streams found in file」，白白等一次往返。"""
    try:
        r = subprocess.run(
            [FFPROBE, "-v", "error", "-select_streams", "a",
             "-show_entries", "stream=index", "-of", "csv=p=0", path],
            capture_output=True, text=True, timeout=30)
        return bool((r.stdout or "").strip())
    except Exception:
        return True     # 探测不出来就按「有」处理，让 ASR 服务自己去报错


def asr_bad_ratio(text: str, duration: float) -> bool:
    """判断已入库的转写文本是不是「复读退化」产物。

    历史包袱：1.0.2 之前的版本把整首音频一次性丢给 ASR（oMLX 默认
    chunk_duration=1200 秒 = 不切），Qwen3-ASR 因此退化成复读机并写进了
    index.db（实测全库 123 首里 29 首中招）。这些记录以前被当成「已转写」
    永远不会重跑，用户看到的就是「同一首歌在 oMLX 能转、在 FXseek 不行」。

    ★ 判定逻辑**只在 asr_desc._looks_degenerate 里有一份**，这里直接调用，
    别再抄一遍 —— 旧版这里和 asr_desc 各写了一份「字/秒 > 8」，而那个阈值
    既漏判韩文/俄文（字数被正则剥成 0），又会误杀正常英文歌。两处一起改才
    不会出现「转写时放行、入库时丢弃」的自相矛盾。

    `duration` 保留在签名里只为兼容调用方，新判定不依赖时长。
    """
    try:
        import asr_desc as _asr
        return _asr._looks_degenerate(text, duration)
    except Exception:
        return False


def asr_status(db_path=None, max_duration=None):
    """音频/视频里，哪些已经转写过、哪些还没有。

    返回 {total, done, pending: [{path, kind, name, bad?}]}。
    「已转写」的判据是 meta 里有非空 asr_text；已经确认过「没有内容」(`asr_none`)
    的算处理过 —— 它不该反复重试。`asr_skip` 则要看**当初是哪条上限挡的**：
    上限被调大（或改成不限）之后，这些文件必须重新排队，判据见 `_asr_skip_lifted`。
    **复读退化的旧记录也要重新排队**（见 asr_bad_ratio），否则永远修不回来。
    """
    db = db_path or DB_PATH
    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT path, kind, meta FROM items"
            " WHERE kind IN ('audio','video') AND chunk_idx<?"
            " GROUP BY path, kind ORDER BY MIN(id)", (AI_TAG_BASE,)).fetchall()
    except sqlite3.OperationalError:
        rows = []
    finally:
        con.close()
    pend = []
    for p, k, mstr in rows:
        try:
            m = json.loads(mstr or "{}")
        except Exception:
            m = {}
        txt = (m.get("asr_text") or "").strip()
        if txt:
            # 老版本写的复读产物：重新排队重转（不计入 done）。
            # 先按字数粗筛，避免给每一首都跑一次 ffprobe（退化样本都几千字）。
            # ★ 粗筛用 asr_desc._plain_len（认所有语种的字母/数字）而不是
            # `[^\u4e00-\u9fffA-Za-z0-9]` 白名单正则 —— 后者会把韩文/俄文/日文
            # 假名整段剥空成 0 字，那些语种的复读产物会永远筛不出来。
            try:
                import asr_desc as _asr
                plain = _asr._plain_len(txt)
            except Exception:
                import re as _re
                plain = len(_re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", txt))
            if plain >= 100:
                dur = float(m.get("duration") or 0) or _probe_duration(p)
                if asr_bad_ratio(txt, dur):
                    pend.append({"path": p, "kind": k,
                                 "name": os.path.basename(p), "bad": True})
            continue
        if m.get("asr_none"):
            continue
        if m.get("asr_skip") and not _asr_skip_lifted(m, max_duration):
            continue
        pend.append({"path": p, "kind": k, "name": os.path.basename(p)})
    return {"total": len(rows), "done": len(rows) - len(pend), "pending": pend}


def transcribe_file(path, asr_cfg, model=None, processor=None, db_path=None):
    """给单个音频/视频补做 ASR 转写，**不重抽帧、不重新解码**。

    与 tag_file 同一套思路：索引期不再做，改成「详情页手动」与「闲置时自动」两处。
    转写文字影响搜索的方式有两种，都要照顾到：

      · 所有音视频：写进 meta.asr_text —— 关键词路（`meta LIKE '%asr_text%'`）与
        L2 的 `_lexical_boost` 都读它。视频的转写**只**走这条路，因为视频帧是
        图片向量，文字拼不进编码。
      · 音频：chunk 0 的文本向量原本就是「音频文件 X，时长 N 秒 + ai_tags + asr_text」
        编码出来的（见 build_index 的编码分支），所以转写完必须把 chunk 0 重编码，
        否则「按歌词搜」在向量这条路上仍然搜不到。重编码只要几十毫秒。

    返回 {ok, text, chars, vectors, seconds} 或 {ok: False, message}。
    """
    import asr_desc as _asr

    db = db_path or DB_PATH
    ext = os.path.splitext(path)[1].lower()
    if ext not in (AUDIO_EXT | VIDEO_EXT):
        return {"ok": False, "message": "只支持音频与视频"}
    if not asr_cfg or not (asr_cfg.get("base_url") and asr_cfg.get("model")):
        return {"ok": False, "message": "请先在设置里配置语音识别服务与模型"}
    if not os.path.exists(path):
        return {"ok": False, "message": "文件不存在"}

    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT kind, chunk_idx, meta, mtime FROM items"
            " WHERE path=? AND chunk_idx<? ORDER BY chunk_idx",
            (path, AI_TAG_BASE)).fetchall()
        if not rows:
            return {"ok": False, "message": "这个素材还没有索引，先刷新索引再转写"}
        kind, mtime = rows[0][0], rows[0][3]
        try:
            m0 = json.loads(rows[0][2] or "{}")
        except Exception:
            m0 = {}

        t0 = time.time()
        max_dur = asr_cfg.get("max_duration") or 0
        if max_dur:
            dur = _probe_duration(path)
            if dur and dur > max_dur:
                # 连「哪条上限挡的」一起记下来，上限放宽后 _asr_pending 才好放行
                _merge_meta(con, path, {"asr_skip": "too_long",
                                        "asr_skip_dur": round(dur, 1)})
                con.commit()
                return {"ok": False,
                        "message": f"时长 {dur/60:.1f} 分钟，超过「设置 → 智能服务 → "
                                   f"音频转写」里的时长上限 {max_dur} 秒。"
                                   f"把那个数字调大、或填 0 表示不限，就能转写。"}
        # 没有音频轨的视频（测试卡、纯画面录像）直接标记，不必去问 ASR 服务
        if not _has_audio_stream(path):
            _merge_meta(con, path, {"asr_none": True})
            con.commit()
            return {"ok": False, "message": "这个文件里没有音频轨，没有可转写的内容"}

        r = _asr.transcribe(asr_cfg["base_url"], asr_cfg.get("api_key", ""),
                            asr_cfg["model"], path,
                            language=asr_cfg.get("language", ""))
        txt = (r.get("text") or "").strip()
        # 最后一道闸：即便分窗后仍然复读（极端音频），也**绝不能把它写进索引**——
        # 一旦落库就会被当成「已转写」，复读文本会污染关键词路和向量，且永不重跑。
        # 宁可报失败让它下次重试。
        dur_chk = _probe_duration(path)
        if txt and asr_bad_ratio(txt, dur_chk):
            return {"ok": False,
                    "message": f"转写结果异常（{len(txt)} 字 / {dur_chk:.0f} 秒，"
                               f"疑似模型复读），已丢弃，稍后重试"}
        if txt:
            asr_meta = {"asr_text": txt, "asr_model": asr_cfg["model"]}
            asr_clear = ("asr_none", "asr_skip", "asr_skip_dur")   # 这次有字了，清掉「没内容/跳过」标记
        else:
            asr_meta = {"asr_none": True}
            asr_clear = ("asr_text", "asr_model", "asr_skip",
                         "asr_skip_dur")   # ★ 清掉上一次的旧文字与跳过痕迹
        _merge_meta(con, path, asr_meta, clear=asr_clear)

        n_vec = 0
        if txt and kind == "audio" and model is not None:
            # 重编码失败不该连累已经转写好的文字：单独兜住，文字照样入库。
            try:
                import embed as we
                import mlx.core as mx
                name = os.path.splitext(os.path.basename(path))[0]
                dur_v = m0.get("duration") or 0
                dur_txt = f"，时长 {dur_v:.0f} 秒" if dur_v else ""
                base_text = f"音频文件 {name}{dur_txt}"
                extra = ""
                if m0.get("ai_tags"):
                    extra = " " + " ".join(m0["ai_tags"])
                vec = we.embed(model, processor, text=base_text + extra + " " + txt,
                               instruction=we.QUERY_INSTRUCTION)
                mx.eval(vec)
                lst = vec.tolist()
                newm = dict(m0)
                newm.update(asr_meta)
                for k in asr_clear:
                    newm.pop(k, None)
                con.execute(
                    "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                    " VALUES (?,?,?,?,?,?,?,?)",
                    (path, kind, rows[0][1], json.dumps(newm, ensure_ascii=False),
                     len(lst), _to_blob(lst), mtime or time.time(), time.time()))
                n_vec = 1
            except Exception as e:
                print(f"  [转写后重编码失败] {path}: {e}")
        con.commit()
        return {"ok": True, "text": txt, "chars": len(txt),
                "vectors": n_vec, "seconds": round(time.time() - t0, 1)}
    except _asr.ASRError as e:
        return {"ok": False, "message": str(e)}
    except Exception as e:
        return {"ok": False, "message": f"{type(e).__name__}: {e}"}
    finally:
        con.close()


def polish_file(path, ai_cfg, model=None, processor=None, db_path=None):
    """用 LLM 把 ASR 原文优化一份，**另存 asr_clean，原文一个字不动**。

    为什么不能覆盖原文：检索有两路 —— 关键词路（`meta LIKE '%asr_text%'`）和
    chunk 0 的向量重编码。用户搜的时候往往**正是照着他听到的那个错字打的**，
    把原文改掉，那些错字就永远搜不到了。所以这里是「加一份」，不是「改一份」。

    音频还要把 chunk 0 重编码（原文 + 优化版一起拼进去），否则「按优化后的词搜」
    在向量这条路上搜不到。视频的转写只走关键词路，不需要重编码。

    返回 {ok, chars, kind, vectors, seconds} 或 {ok: False, message}。
    """
    import asr_polish as _pol

    db = db_path or DB_PATH
    if not ai_cfg or not (ai_cfg.get("base_url") and ai_cfg.get("model")):
        return {"ok": False, "message": "请先在设置里配置智能服务的地址与模型"}
    if not os.path.exists(path):
        return {"ok": False, "message": "文件不存在"}

    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT kind, chunk_idx, meta, mtime FROM items"
            " WHERE path=? AND chunk_idx<? ORDER BY chunk_idx",
            (path, AI_TAG_BASE)).fetchall()
        if not rows:
            return {"ok": False, "message": "这个素材还没有索引，先刷新索引"}
        kind, mtime = rows[0][0], rows[0][3]
        try:
            m0 = json.loads(rows[0][2] or "{}")
        except Exception:
            m0 = {}
        src = m0.get("asr_text") or ""
        if not src.strip():
            return {"ok": False, "message": "这个素材还没有转写文字，先转写再优化"}

        t0 = time.time()
        r = _pol.polish_full(src, path, ai_cfg["base_url"],
                             ai_cfg.get("api_key", ""), ai_cfg["model"])
        if not r.get("ok"):
            return {"ok": False, "message": r.get("message") or "优化失败"}
        txt = r["text"]
        pmeta = {"asr_clean": txt,
                 "asr_clean_model": ai_cfg["model"],
                 "asr_clean_kind": r.get("kind") or _pol.MONOLOGUE,
                 "asr_clean_at": time.time()}
        _merge_meta(con, path, pmeta)

        n_vec = 0
        if kind == "audio" and model is not None:
            # 重编码失败不该连累已经写好的优化文字：单独兜住。
            try:
                import embed as we
                import mlx.core as mx
                name = os.path.splitext(os.path.basename(path))[0]
                dur_v = m0.get("duration") or 0
                dur_txt = f"，时长 {dur_v:.0f} 秒" if dur_v else ""
                extra = ""
                if m0.get("ai_tags"):
                    extra = " " + " ".join(m0["ai_tags"])
                # 原文在前、优化版在后：两边都要能被向量命中
                body = f"{src} {txt}"
                vec = we.embed(model, processor,
                               text=f"音频文件 {name}{dur_txt}" + extra + " " + body,
                               instruction=we.QUERY_INSTRUCTION)
                mx.eval(vec)
                lst = vec.tolist()
                newm = dict(m0)
                newm.update(pmeta)
                con.execute(
                    "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                    " VALUES (?,?,?,?,?,?,?,?)",
                    (path, kind, rows[0][1], json.dumps(newm, ensure_ascii=False),
                     len(lst), _to_blob(lst), mtime or time.time(), time.time()))
                n_vec = 1
            except Exception as e:
                print(f"  [优化后重编码失败] {path}: {e}")
        con.commit()
        return {"ok": True, "chars": len(txt), "kind": pmeta["asr_clean_kind"],
                "vectors": n_vec, "seconds": round(time.time() - t0, 1)}
    except Exception as e:
        return {"ok": False, "message": f"{type(e).__name__}: {e}"}
    finally:
        con.close()


def unpolish_file(path, model=None, processor=None, db_path=None):
    """丢掉优化版，回到只看原文。原文一直都在，所以元数据是纯删除。

    但音频**光删元数据不够**：polish_file 当初把「原文 + 优化版」拼进了 chunk 0
    的向量，不重编码的话，优化版里的词还能从向量这条路搜到这个文件 ——
    用户点了「还原」，结果搜优化版才有的词还能搜出来，那就是还原得不干净。
    所以音频要把 chunk 0 按**只剩原文**重编码一遍（视频的转写不走向量，无需处理）。
    """
    db = db_path or DB_PATH
    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT kind, chunk_idx, meta, mtime FROM items"
            " WHERE path=? AND chunk_idx<? ORDER BY chunk_idx",
            (path, AI_TAG_BASE)).fetchall()
        if not rows:
            return {"ok": False, "message": "这个素材还没有索引"}
        kind, mtime = rows[0][0], rows[0][3]
        try:
            m0 = json.loads(rows[0][2] or "{}")
        except Exception:
            m0 = {}
        had = bool((m0.get("asr_clean") or "").strip())

        _merge_meta(con, path,
                    {}, clear=("asr_clean", "asr_clean_model",
                               "asr_clean_kind", "asr_clean_at"))

        n_vec = 0
        if had and kind == "audio" and model is not None:
            try:
                import embed as we
                import mlx.core as mx
                src = (m0.get("asr_text") or "").strip()
                if src:
                    name = os.path.splitext(os.path.basename(path))[0]
                    dur_v = m0.get("duration") or 0
                    dur_txt = f"，时长 {dur_v:.0f} 秒" if dur_v else ""
                    extra = ""
                    if m0.get("ai_tags"):
                        extra = " " + " ".join(m0["ai_tags"])
                    vec = we.embed(model, processor,
                                   text=f"音频文件 {name}{dur_txt}" + extra + " " + src,
                                   instruction=we.QUERY_INSTRUCTION)
                    mx.eval(vec)
                    lst = vec.tolist()
                    newm = dict(m0)
                    for k in ("asr_clean", "asr_clean_model",
                              "asr_clean_kind", "asr_clean_at"):
                        newm.pop(k, None)
                    con.execute(
                        "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                        " VALUES (?,?,?,?,?,?,?,?)",
                        (path, kind, rows[0][1], json.dumps(newm, ensure_ascii=False),
                         len(lst), _to_blob(lst), mtime or time.time(), time.time()))
                    n_vec = 1
            except Exception as e:
                # 向量没救回来不算致命：元数据已经干净了，下次刷新索引会自愈
                print(f"  [还原后重编码失败] {path}: {e}")
        con.commit()
        return {"ok": True, "vectors": n_vec}
    finally:
        con.close()


def polish_progress(db_path=None):
    """统计「有转写文字」和「已优化」的条数，给「一键优化全部」显示进度用。"""
    db = db_path or DB_PATH
    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT meta FROM items WHERE chunk_idx=0 AND kind IN ('audio','video')"
        ).fetchall()
        total = done = 0
        for (ms,) in rows:
            try:
                m = json.loads(ms or "{}")
            except Exception:
                m = {}
            if (m.get("asr_text") or "").strip():
                total += 1
                if (m.get("asr_clean") or "").strip():
                    done += 1
        return {"total": total, "done": done}
    finally:
        con.close()


def polish_pending(db_path=None):
    """还没优化过、但有转写文字可优化的素材路径。"""
    db = db_path or DB_PATH
    con = sqlite3.connect(db)
    try:
        rows = con.execute(
            "SELECT path, meta FROM items WHERE chunk_idx=0 AND kind IN ('audio','video')"
        ).fetchall()
        out = []
        for path, ms in rows:
            try:
                m = json.loads(ms or "{}")
            except Exception:
                m = {}
            if (m.get("asr_text") or "").strip() and not (m.get("asr_clean") or "").strip():
                out.append(path)
        return out
    finally:
        con.close()


# --------------------------------------------------------------------------
# 数据库
# --------------------------------------------------------------------------
def _init_db(db_path=DB_PATH):
    os.makedirs(os.path.dirname(db_path), exist_ok=True)
    con = sqlite3.connect(db_path)
    con.execute("""
        CREATE TABLE IF NOT EXISTS items (
            id          INTEGER PRIMARY KEY AUTOINCREMENT,
            path        TEXT,
            kind        TEXT,
            chunk_idx   INTEGER DEFAULT 0,
            meta        TEXT,
            dim         INTEGER,
            vec         BLOB,
            mtime       REAL,
            indexed_at  REAL,
            UNIQUE(path, chunk_idx)
        )
    """)
    con.execute("CREATE INDEX IF NOT EXISTS idx_path ON items(path)")
    con.execute("CREATE INDEX IF NOT EXISTS idx_kind ON items(kind)")
    con.commit()
    _migrate_db(con)
    return con


def _migrate_db(con):
    """旧库迁移：把 path UNIQUE 改为 (path, chunk_idx) UNIQUE。

    旧版本用 path 做唯一键，导致同一文件的多帧/多分块互相覆盖。
    """
    try:
        row = con.execute(
            "SELECT sql FROM sqlite_master WHERE type='table' AND name='items'"
        ).fetchone()
        if not row or not row[0]:
            return False
        compact = row[0].replace(" ", "").replace("\n", "")
        if "(path,chunk_idx)" in compact.replace("UNIQUE(path,chunk_idx)", "(path,chunk_idx)"):
            return False
        con.execute("ALTER TABLE items RENAME TO items_old")
        con.execute("""
            CREATE TABLE items (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                path TEXT, kind TEXT, chunk_idx INTEGER DEFAULT 0,
                meta TEXT, dim INTEGER, vec BLOB, mtime REAL, indexed_at REAL,
                UNIQUE(path, chunk_idx)
            )
        """)
        con.execute("""
            INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)
            SELECT path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at FROM items_old
        """)
        con.execute("DROP TABLE items_old")
        con.execute("CREATE INDEX IF NOT EXISTS idx_path ON items(path)")
        con.execute("CREATE INDEX IF NOT EXISTS idx_kind ON items(kind)")
        con.commit()
        print("[migrate] 索引库已升级为 (path, chunk_idx) 唯一键")
        return True
    except Exception as e:
        print(f"[migrate] 迁移跳过: {e}")
        return False


def _to_blob(vec):
    return struct.pack(f"<{len(vec)}f", *[float(x) for x in vec])


def _from_blob(blob):
    n = len(blob) // 4
    return list(struct.unpack(f"<{n}f", blob))


def _file_id(path: str) -> str:
    return hashlib.sha1(path.encode("utf-8")).hexdigest()[:16]


# --------------------------------------------------------------------------
# 文件解析：抽出「可 embedding 的单元」
# --------------------------------------------------------------------------
def _probe_duration(path: str) -> float:
    """用 ffprobe 取媒体时长（秒），失败返回 0。"""
    try:
        out = subprocess.run(
            [FFPROBE, "-v", "error", "-show_entries", "format=duration",
             "-of", "default=nw=1:nk=1", path],
            capture_output=True, text=True, timeout=30)
        return float(out.stdout.strip() or 0)
    except Exception:
        return 0.0


def _asr_skip_lifted(m: dict, max_duration) -> bool:
    """meta 里的「跳过转写」标记是不是**因为时长上限被放宽而作废了**。

    ★ 为什么要有这个函数：判断「该不该重试」的地方有**两处** ——
      `_asr_pending()`（增量索引时决定跳不跳过）和 `asr_status()`（界面上的
      待转写清单）。以前两处各写各的 `if m.get("asr_skip"): skip`，于是
      「用户把上限从 600 调到 0」这件事它们都看不见，那些 30~60 分钟的
      专辑/演唱会 FLAC 会**永远停在跳过状态**，一次都不会再试。
      抽成一份共用判断，以后改规则只改这里。

    只有 `too_long` 这一种理由是可以被放宽救回来的；`asr_none`（确认没语音）
    之类的标记与上限无关，永远不该重试。
    """
    if m.get("asr_skip") != "too_long":
        return False
    try:
        max_dur = int(max_duration or 0)
    except Exception:
        max_dur = 0
    if not max_dur:
        return True                       # 已改成「不限」→ 放回来重试
    try:
        dur = float(m.get("duration") or 0)
    except Exception:
        dur = 0.0
    if not dur:
        return True                       # 没记时长，交给下游自己再量一次
    return dur <= max_dur                 # 上限调高到够得着了 → 重试


def _asr_pending(cur, path: str, asr) -> bool:
    """ASR 开着、但这个音视频文件还没有转写文本 → 需要重跑。

    历史遗留：增量索引遇到 mtime 未变就整个跳过，导致开启 ASR 之前入库的
    歌曲/视频从来没被转写过（meta 里没有 asr_text），歌词也就搜不到。
    这里让这类文件不参与跳过，落回正常索引分支复用全部 ASR + 编码逻辑。
    已确认过「转写为空」的文件、以及「仍超时长上限」的文件不再反复重试。
    """
    if not asr:
        return False
    if os.path.splitext(path)[1].lower() not in (AUDIO_EXT | VIDEO_EXT):
        return False
    try:
        row = cur.execute("SELECT meta FROM items WHERE path=? LIMIT 1",
                          (path,)).fetchone()
    except Exception:
        return False
    if not row:
        return False
    try:
        m = json.loads(row[0] or "{}")
    except Exception:
        m = {}
    if (m.get("asr_text") or "").strip():
        return False
    if m.get("asr_none"):
        return False
    if m.get("asr_skip"):
        return _asr_skip_lifted(m, (asr or {}).get("max_duration"))
    return True


def _tags_pending(cur, path: str) -> bool:
    """音频：这一条是「标签入库」之前编的 → 需要重编一次文本侧。

    判定只看库里有没有 tags_checked 标记，**不去碰文件系统**（不额外跑 ffprobe）：
      有标记 = 已经按新模板编过（不管有没有标签），跳过；
      没标记 = 老记录，放回正常分支重编一次 —— 这就是老库的平滑迁移通道。
    只认 AUDIO_EXT：视频重编要重新抽帧，代价太大，不在这次迁移范围里。
    """
    if os.path.splitext(path)[1].lower() not in AUDIO_EXT:
        return False
    try:
        row = cur.execute("SELECT meta FROM items WHERE path=? AND chunk_idx=0 LIMIT 1",
                          (path,)).fetchone()
    except Exception:
        return False
    if not row:
        return False
    try:
        m = json.loads(row[0] or "{}")
    except Exception:
        return False
    return not m.get("tags_checked")


# 抽帧密度模式：自动策略的缩放系数
DENSITY_SCALE = {
    "dense": 0.5,     # 密集：间隔减半（更准，更耗资源）
    "auto": 1.0,      # 自适应（默认）
    "sparse": 2.0,    # 稀疏：间隔翻倍（更省资源）
}


def frame_plan(duration: float, max_frames: int = None,
               mode: str = "auto", custom_gap: float = None):
    """返回 (帧间隔, 帧数上限, 抽帧时间点列表)。

    自适应策略：短视频密抽（任意画面都能命中），长视频疏抽（控制资源）。
    mode: auto(自适应) | dense(密集) | sparse(稀疏) | custom(固定间隔)
    """
    dur = float(duration or 0)
    if dur <= 0:
        return 3.0, 4, [0.0, 1.0, 3.0, 5.0]

    if mode == "custom" and custom_gap and custom_gap > 0:
        gap, cap = float(custom_gap), VIDEO_MAX_FRAMES
    else:
        gap, cap = 3.0, 8
        for limit, g, c in VIDEO_FRAME_TIERS:
            if dur <= limit:
                gap, cap = g, c
                break
        scale = DENSITY_SCALE.get(mode, 1.0)
        gap = max(0.1, gap * scale)
        cap = int(cap / scale) if scale > 0 else cap

    if max_frames and max_frames > 0:
        cap = min(cap, max_frames)
    n = max(1, min(cap, int(dur / gap) + 1))
    times = [round(dur * (i + 0.5) / n, 2) for i in range(n)]
    return gap, cap, times


# 运行时抽帧配置（由 app 层根据设置注入）
FRAME_OPTS = {"mode": "auto", "custom_gap": None, "max_frames": None}


def _video_frames(path: str, max_frames=None):
    """用 ffmpeg 抽帧（自适应密度），返回 [(PIL.Image, timestamp), ...]"""
    from PIL import Image
    frames = []
    dur = _probe_duration(path)
    _gap, _cap, times = frame_plan(dur,
                                   max_frames=max_frames or FRAME_OPTS.get("max_frames"),
                                   mode=FRAME_OPTS.get("mode", "auto"),
                                   custom_gap=FRAME_OPTS.get("custom_gap"))
    print(f"  [抽帧] {os.path.basename(path)} 时长 {dur:.1f}s → {len(times)} 帧 (间隔 {_gap}s, 模式 {FRAME_OPTS.get('mode','auto')})")

    with tempfile.TemporaryDirectory() as td:
        for i, t in enumerate(times):
            fp = os.path.join(td, f"f{i}.jpg")
            try:
                subprocess.run(
                    [FFMPEG, "-ss", str(t), "-i", path,
                     "-frames:v", "1", "-q:v", "3", "-y", fp],
                    capture_output=True, timeout=60)
                if os.path.exists(fp):
                    frames.append((Image.open(fp).convert("RGB").copy(), round(t, 2)))
            except Exception:
                continue
    return frames


# 视频关键帧九宫格：一次调用让视觉模型「看完整段视频」
VIDEO_SHEET_PROMPT = (
    "这是同一段视频按时间顺序抽取的若干关键帧拼成的网格图（从左到右、从上到下）。"
    "你是素材库打标助手，请观察整段视频内容，生成用于语义检索的描述和标签。\n\n"
    "严格按以下两行输出，不要加任何解释、前后缀或多余文字：\n\n"
    "描述：一句简洁中文，概括这段视频的主体、人物、动作与环境。\n\n"
    "关键词：逗号分隔 8~12 个短词（每个 2~6 字），要具体、有区分度，覆盖：人物/主体"
    "是什么、动作、环境、关键物体、整体风格。禁止用「视频」「场景」「画面」「内容」"
    "这类泛词；同一概念只留最具体的一个词；无法识别时关键词写「无」。"
)


def _video_contact_sheet(path: str, cells: int = 9, cell: int = 320):
    """把视频的关键帧拼成一张网格图，返回 PIL.Image（失败返回 None）。

    为什么要拼图：逐帧打标＝每帧一次大模型调用，一分钟的视频就是几十次调用、
    几十分钟；拼成一张九宫格只要 **一次** 调用，代价与视频长度无关。
    """
    from PIL import Image
    try:
        frames = _video_frames(path, max_frames=cells)
    except Exception as e:
        print(f"  [拼图抽帧失败] {os.path.basename(path)}: {e}")
        return None
    if not frames:
        return None
    step = max(1, len(frames) // cells)
    sel = frames[::step][:cells]
    if not sel:
        return None
    cols = max(1, int(len(sel) ** 0.5 + 0.999))
    rows = max(1, (len(sel) + cols - 1) // cols)
    gap = 6
    W = cols * cell + (cols + 1) * gap
    H = rows * cell + (rows + 1) * gap
    sheet = Image.new("RGB", (W, H), (255, 255, 255))
    for i, (im, _ts) in enumerate(sel):
        r, c = divmod(i, cols)
        t = im.copy()
        t.thumbnail((cell, cell))
        x = gap + c * (cell + gap) + (cell - t.width) // 2
        y = gap + r * (cell + gap) + (cell - t.height) // 2
        sheet.paste(t, (x, y))
    print(f"  [拼图] {os.path.basename(path)} 用 {len(sel)} 帧拼成 {cols}x{rows} 网格 "
          f"({W}x{H})，只需 1 次视觉模型调用")
    return sheet


def _pdf_text(path: str):
    """优先用 pdfplumber 提取 PDF 文本，失败时回退到 textutil（对 doc/docx/rtf 有效）。"""
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        # 优先用 pdfplumber（更好的 PDF 文本提取）
        try:
            import pdfplumber
            text_parts = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            if text_parts:
                return "\n".join(text_parts)
        except Exception as e:
            print(f"  [PDF文本提取-pdfplumber失败] {path}: {e}")
    
    # 回退到 macOS 自带 textutil（对 doc/docx/rtf/wps 有效）
    try:
        out = subprocess.run(["/usr/bin/textutil", "-convert", "txt",
                              "-stdout", path],
                             capture_output=True, text=True, timeout=60)
        if out.returncode == 0 and out.stdout.strip():
            return out.stdout
        if out.returncode == 0:
            # 退出码 0 但无输出 = 文档本身为空（不是解析失败），无需报警
            print(f"  [文档为空] {os.path.basename(path)}")
        else:
            print(f"  [文本提取失败] textutil 返回码={out.returncode}, "
                  f"stderr={out.stderr[:200]}: {path}")
    except Exception as e:
        print(f"  [文本提取异常] {path}: {e}")
    return ""


def _is_probably_text(path: str, sniff: int = 4096) -> bool:
    try:
        with open(path, "rb") as f:
            raw = f.read(sniff)
        if b"\x00" in raw:
            return False
        raw.decode("utf-8", errors="strict")
        return True
    except Exception:
        return False


def _is_ole2(path: str) -> bool:
    """OLE2 复合文档（D0CF11E0），即老式 .doc/.xls/.ppt 以及 WPS 的 .wps/.et/.dps。"""
    try:
        with open(path, "rb") as f:
            return f.read(8) == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    except Exception:
        return False


def _ole_biff_text(path: str) -> str:
    """读 OLE2 里的老式 BIFF 工作簿（.xls / WPS .et）。

    .et 是 WPS 表格的老二进制格式：OLE2 容器里放一个 `Workbook` 流，
    内容是 BIFF 记录，用 xlrd 直接解析这条流即可。
    """
    try:
        import olefile
        import xlrd
    except Exception as e:
        print(f"  [表格解析] 缺少 olefile/xlrd，跳过 {path}: {e}")
        return ""

    try:
        ole = olefile.OleFileIO(path)
    except Exception as e:
        print(f"  [OLE2打开失败] {path}: {e}")
        return ""

    try:
        names = {"/".join(s) for s in ole.listdir()}
        # WPS .et 与老 .xls 都用 Workbook；少数用 Book
        stream = None
        for cand in ("Workbook", "Book"):
            if cand in names:
                stream = cand
                break
        if not stream:
            print(f"  [表格解析] {path} 内无 Workbook 流，实际流: {sorted(names)}")
            return ""
        data = ole.openstream(stream).read()
    except Exception as e:
        print(f"  [OLE2读流失败] {path}: {e}")
        return ""
    finally:
        try:
            ole.close()
        except Exception:
            pass

    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception as e:
        print(f"  [BIFF解析失败] {path}: {e}")
        return ""

    parts = []
    for sh in book.sheets():
        rows_txt = []
        for r in range(min(sh.nrows, 5000)):     # 单表上限 5000 行
            cells = []
            for c in range(sh.ncols):
                try:
                    v = sh.cell_value(r, c)
                except Exception:
                    v = ""
                if v is None:
                    cells.append("")
                elif isinstance(v, float) and v.is_integer():
                    cells.append(str(int(v)))    # 128000.0 → 128000
                else:
                    cells.append(str(v).strip())
            line = "\t".join(cells).strip()
            if line.replace("\t", ""):
                rows_txt.append(line)
        if sh.nrows > 5000:
            rows_txt.append("……（表格过长，已截断）")
        if rows_txt:
            parts.append(f"【工作表：{sh.name}】\n" + "\n".join(rows_txt))
    return "\n\n".join(parts)


def _xlsx_text(path: str) -> str:
    """用 openpyxl 读表格单元格文本。

    按「行」拼接（制表符分隔），并带上工作表名与表头，让「某列叫什么」也能被检索到。
    公式取公式串本身；日期等类型直接 str()。

    注意：openpyxl 会按扩展名做白名单校验，`.et`（WPS 表格）虽然内部就是 xlsx 容器
    也会被拒。因此非标准后缀先复制成临时 .xlsx 再读。
    """
    try:
        import openpyxl
    except Exception as e:
        print(f"  [表格解析] 缺少 openpyxl，跳过 {path}: {e}")
        return ""

    ext = Path(path).suffix.lower()
    tmp = None
    target = path
    if ext not in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        # .et 等：拷成合法后缀再交给 openpyxl
        try:
            import shutil
            fd, tmp = tempfile.mkstemp(suffix=".xlsx")
            os.close(fd)
            shutil.copyfile(path, tmp)
            target = tmp
        except Exception as e:
            print(f"  [表格解析-临时副本失败] {path}: {e}")
            return ""

    try:
        # read_only + data_only=False：大表不爆内存，保留公式文本
        wb = openpyxl.load_workbook(target, read_only=True, data_only=False)
    except Exception as e:
        print(f"  [表格打开失败] {path}: {e}")
        return ""

    parts = []
    try:
        for ws in wb.worksheets:
            rows_txt = []
            for r_i, row in enumerate(ws.iter_rows(values_only=True)):
                if r_i > 5000:          # 单表最多扫 5000 行，防超大表拖垮索引
                    rows_txt.append("……（表格过长，已截断）")
                    break
                cells = [("" if c is None else str(c).strip()) for c in row]
                line = "\t".join(cells).strip()
                if line.replace("\t", ""):
                    rows_txt.append(line)
            if rows_txt:
                parts.append(f"【工作表：{ws.title}】\n" + "\n".join(rows_txt))
    finally:
        try:
            wb.close()
        except Exception:
            pass
        if tmp:
            try:
                os.remove(tmp)
            except Exception:
                pass
    return "\n\n".join(parts)


def _pptx_text(path: str) -> str:
    """用 python-pptx 读幻灯片文本框。

    .dps（WPS 演示）本质也是 OOXML zip 包，python-pptx 多数情况能直接读；
    读不了时回退到解 zip 抽 <a:t> 文本节点。
    """
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        texts.append(t)
                # 表格形状
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells]
                        line = "\t".join(cells).strip()
                        if line.replace("\t", ""):
                            texts.append(line)
            if texts:
                parts.append(f"【第 {i} 页】\n" + "\n".join(texts))
        if parts:
            return "\n\n".join(parts)
    except Exception as e:
        print(f"  [PPT解析-python-pptx失败] {path}: {e}")

    # 回退：直接解 zip 抽 <a:t> 文本（.dps / 异常 pptx）
    try:
        import zipfile
        import re as _re
        chunks = []
        with zipfile.ZipFile(path) as z:
            names = [n for n in z.namelist()
                     if _re.match(r"ppt/slides/slide\d+\.xml$", n)]
            names.sort(key=lambda n: int(_re.findall(r"\d+", n)[-1]))
            for n in names:
                xml = z.read(n).decode("utf-8", errors="ignore")
                txts = _re.findall(r"<a:t>(.*?)</a:t>", xml, _re.S)
                txts = [t.strip() for t in txts if t.strip()]
                if txts:
                    chunks.append("\n".join(txts))
        if chunks:
            return "\n\n".join(chunks)
    except Exception as e:
        print(f"  [PPT解析-zip回退失败] {path}: {e}")
    return ""


def _strip_html(html: str) -> str:
    """HTML 只留可见文本。

    空模板（`<!DOCTYPE html>…<body></body>`）如果原样入库，标签本身就是
    「低信息但通用」的文本，会像空文档一样对任何查询给 0.5+ 的相似度。
    """
    import re as _re
    s = _re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", html)
    s = _re.sub(r"(?s)<[^>]+>", " ", s)
    for a, b in (("&nbsp;", " "), ("&lt;", "<"), ("&gt;", ">"),
                 ("&quot;", '"'), ("&amp;", "&")):
        s = s.replace(a, b)
    return " ".join(s.split())


def _read_text(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _pdf_text(path)
    if ext in (".doc", ".docx", ".rtf", ".wps"):
        return _pdf_text(path)
    if ext == ".xlsx":
        return _xlsx_text(path)
    if ext == ".et":
        # WPS 表格：可能是老式 OLE2/BIFF（真 .et），也可能是 xlsx 换壳，两种都试
        if _is_ole2(path):
            return _ole_biff_text(path)
        return _xlsx_text(path)
    if ext in (".xlsm",):
        return _xlsx_text(path)
    if ext == ".pptx":
        return _pptx_text(path)
    if ext == ".dps":
        # WPS 演示：老式 OLE2/PowerPoint 二进制读不了内容（无开源解析器），
        # 但仍要收录，否则用户在库里看不到该文件；用文件名兜底见 iter_units。
        return "" if _is_ole2(path) else _pptx_text(path)
    if _is_probably_text(path):
        for enc in ("utf-8", "gbk", "latin-1"):
            try:
                raw = Path(path).read_text(encoding=enc)
                break
            except Exception:
                continue
        else:
            return ""
        if ext in (".html", ".htm"):
            return _strip_html(raw)
        return raw
    return ""


def _chunk(text: str, size=DOC_CHUNK_CHARS, overlap=DOC_CHUNK_OVERLAP):
    text = " ".join(text.split())
    if not text:
        return []
    if len(text) <= size:
        return [text]
    chunks, i = [], 0
    while i < len(text):
        chunks.append(text[i:i + size])
        i += size - overlap
    return chunks


def iter_units(path: str):
    """把一个文件拆成可索引单元，yield (kind, chunk_idx, meta, text, image)。"""
    ext = Path(path).suffix.lower()
    if ext in IMAGE_EXT:
        from PIL import Image
        try:
            img = Image.open(path).convert("RGB")
            yield ("image", 0, {}, None, img, )
        except Exception as e:
            print(f"  [跳过] 图片打开失败 {path}: {e}")
    elif ext in VIDEO_EXT:
        for idx, (img, ts) in enumerate(_video_frames(path)):
            yield ("video", idx, {"timestamp": ts}, None, img)
    elif ext in AUDIO_EXT:
        # 音频：WeMM 不处理音频内容，用文件名 + 容器标签构造文本向量
        #（这样能被「文件名 / 关键词 / 歌手 / 专辑 / 流派」搜到）。
        name = os.path.splitext(os.path.basename(path))[0]
        # ★ 一次 ffprobe 同时拿时长与标签 —— 以前只取时长，tags 白扔了。
        mm = _probe_media_meta_raw(path)
        tg = mm.get("tags") or {}
        dur = float(mm.get("duration") or 0) or _probe_duration(path)
        dur_txt = f"，时长 {dur:.0f} 秒" if dur > 0 else ""
        parts = []
        if tg.get("title") and tg["title"] != name:
            parts.append(f"曲名 {tg['title']}")      # 曲名和文件名重复时不用再写一遍
        if tg.get("artist"):
            parts.append(f"歌手 {tg['artist']}")
        if tg.get("album"):
            parts.append(f"专辑 {tg['album']}")
        if tg.get("genre"):
            parts.append(f"流派 {tg['genre']}")
        if tg.get("year"):
            parts.append(f"年份 {tg['year']}")
        if tg.get("composer"):
            parts.append(f"作曲 {tg['composer']}")
        tag_txt = ("，" + "，".join(parts)) if parts else ""
        text = f"音频文件 {name}{dur_txt}{tag_txt}"
        yield ("audio", 0, {"duration": round(dur, 1), "tags": tg,
                            "tags_checked": True}, text, None)
    elif ext in DOC_EXT:
        text = _read_text(path)
        chunks = _chunk(text)
        no_text = False
        if not chunks:
            # 提取不到正文（扫描版 PDF、老式二进制 .dps、空 HTML 模板等）也要收录，
            # 否则文件在库里直接消失。用文件名兜底，至少能被文件名检索到。
            name = os.path.splitext(os.path.basename(path))[0]
            chunks = [f"文档 {name}（未能提取正文内容）"]
            no_text = True
        doc_len = len(" ".join(text.split()))
        for i, ch in enumerate(chunks):
            # text_len 是**整篇**正文长度（不是当前分块），供排序降权用：正文极短的
            # 文档语义向量等于「一个空文档」，对任何查询都有 0.5~0.64 相似度
            # （见 _low_info_doc）。
            yield ("document", i, {"text_len": doc_len, "no_text": no_text}, ch, None)


# --------------------------------------------------------------------------
# 索引
# --------------------------------------------------------------------------
def build_index(roots, db_path=DB_PATH, force=False, model_path=None, progress=None,
                ai_config=None, asr_config=None, kinds=None):
    """建立/更新索引。

    progress: 可选回调 fn(dict) —— 用于把实时进度上报给 UI。
      字段：phase(scanning/loading/encoding/done/error), total, current,
            indexed, skipped, units, message
    ai_config: 可选 dict —— AI 描述生成配置（enabled/base_url/api_key/model/prompt）。
               开启后为图片生成描述与标签写入 meta，提升语义检索泛化能力。
    kinds: 可选 list —— 只索引这些类型（image/video/audio/document）。
       None 或空 = 全部类型。用于「刷新索引」跟随当前分类：
       选中「图片」时只扫图片，选中「全部」才全局扫。
       不影响 purge_stale —— 它只删「磁盘上已不存在」与垃圾文件，
       因此限类型扫描不会误删其他类型的索引记录。
    """
    import mlx.core as mx
    import embed as we

    # ★ 索引期**不再**做 AI 打标与 ASR 转写。
    #   这两件事都是「一个素材一次模型调用」，一旦用户加进来一个几千张图的文件夹，
    #   索引就会在没有任何反馈的情况下卡上几小时。现在改由两条更可控的路完成：
    #     ① 详情页的「立即打标签 / 立即转写」按钮（用户自己决定何时花这份算力）；
    #     ② 「闲置时自动处理」守护线程（见 app.py 的 _idle_loop），一次一个，慢慢补齐。
    #   所以最终仍然是「全部文件都有标签、都转写过」，只是把算力摊到了闲时。
    #   ai_config / asr_config 参数保留仅为兼容旧调用方。
    ai = None
    asr = None
    n_ai = 0
    n_asr = 0

    def report(**kw):
        if progress:
            try: progress(kw)
            except Exception: pass

    model_path = model_path or we.DEFAULT_MODEL
    report(phase="loading", message="加载模型…", total=0, current=0,
           indexed=0, skipped=0, units=0)
    print(f"加载模型: {model_path}")
    model, processor = we.load_model(model_path)
    print("模型加载完成 ✅")

    con = _init_db(db_path)
    cur = con.cursor()

    report(phase="scanning", message="扫描文件…")
    # 类型白名单：kinds 为空 → 全部四类；否则只扫指定类别。
    _EXT_OF_KIND = {"image": IMAGE_EXT, "video": VIDEO_EXT,
                    "audio": AUDIO_EXT, "document": DOC_EXT}
    _wanted = [k for k in (kinds or []) if k in _EXT_OF_KIND]
    if _wanted:
        scan_ext = set()
        for k in _wanted:
            scan_ext |= _EXT_OF_KIND[k]
        print(f"仅扫描类型：{'/'.join(_wanted)}")
    else:
        scan_ext = IMAGE_EXT | VIDEO_EXT | AUDIO_EXT | DOC_EXT
    files = []
    for root in roots:
        rp = Path(os.path.realpath(os.path.expanduser(root)))
        if rp.is_file():
            if (not is_junk_file(rp.name)
                    and os.path.splitext(rp.name)[1].lower() in scan_ext):
                files.append(str(rp))
        elif rp.is_dir():
            # 用 os.walk 而不是 rglob：rglob 会一路走进点目录，
            # 于是 ".Trashes/正常名.jpg" 这种也会被当成素材收进来。
            for dp, dns, fns in os.walk(rp):
                dns[:] = [d for d in dns if not d.startswith(".")]
                for fn in fns:
                    if is_junk_file(fn):
                        continue
                    if os.path.splitext(fn)[1].lower() in scan_ext:
                        files.append(os.path.join(dp, fn))
    total = len(files)
    print(f"扫描到 {total} 个候选文件")
    report(phase="encoding", total=total, current=0, message=f"共 {total} 个文件")

    n_new = n_skip = n_unit = 0
    n_fail = 0          # 扫到了但解不出任何单元（坏图/不支持的编码）
    n_meta = 0          # 文件没变，只是补元数据重编了文本侧（不算新增）
    t0 = time.time()
    for fi, path in enumerate(files, 1):
        try:
            mtime = os.path.getmtime(path)
        except OSError:
            continue
        row = cur.execute("SELECT mtime FROM items WHERE path=? LIMIT 1", (path,)).fetchone()
        unchanged = bool(row and not force and abs(row[0] - mtime) < 1e-6)
        asr_need = unchanged and _asr_pending(cur, path, asr)
        # 老记录还没写过 tags_checked → 说明它是「标签入库」之前编的，
        # 放回正常分支重编一次（只动文本侧，很便宜），把歌手/专辑拼进向量。
        meta_need = (unchanged and not asr_need and _tags_pending(cur, path))
        if unchanged and not asr_need and not meta_need:
            # 跳过的音频/视频：若指纹缺失则补建（增量索引不会重复执行新文件分支）
            try:
                ext_fp = os.path.splitext(path)[1].lower()
                if ext_fp in (AUDIO_EXT | VIDEO_EXT):
                    import fingerprint as _fp
                    if not _fp.has_fingerprint(path):
                        _fp.index_file(path, "video" if ext_fp in VIDEO_EXT else "audio", mtime)
            except Exception as e:
                print(f"  [指纹补建失败] {path}: {e}")
            n_skip += 1
            if fi % 5 == 0 or fi == total:
                report(phase="encoding", total=total, current=fi,
                       indexed=n_new, skipped=n_skip, units=n_unit,
                       message=os.path.basename(path))
            continue
        if meta_need:
            n_meta += 1          # 只是补元数据重编，不算「新增文件」

        # ★ 重建之前先把旧 meta 读出来当底稿。
        #   增量刷新时 ai/asr 配置通常是关着的（do_index 传 None），若直接
        #   INSERT 覆盖，会把上一次辛苦跑出来的 asr_text（歌词）和 ai_tags 弄丢 ——
        #   以前「mtime 变了就重索引」的图片/音频就是这么悄悄掉标签的。
        prev_meta = {}
        try:
            _pr = cur.execute(
                "SELECT meta FROM items WHERE path=? AND chunk_idx=0 LIMIT 1",
                (path,)).fetchone()
            if _pr and _pr[0]:
                prev_meta = json.loads(_pr[0]) or {}
        except Exception:
            prev_meta = {}

        cur.execute("DELETE FROM items WHERE path=?", (path,))
        try:
            units = list(iter_units(path))
        except Exception as e:
            print(f"  [错误] {path}: {e}")
            continue

        # AI 描述生成（仅图片，且开启配置时；纯色/低信息图跳过）
        ai_meta = {}
        if ai:
            try:
                from PIL import Image as _PILImage, ImageStat as _ImageStat
                import ai_desc as _ai
                ext_l = os.path.splitext(path)[1].lower()
                skip_ai = False
                # ---- 打标节制策略 ----
                strategy = ai.get("strategy", "smart")
                if strategy == "off":
                    skip_ai = True
                elif ai.get("max_per_batch") and n_ai >= int(ai["max_per_batch"]):
                    skip_ai = True                      # 达到单批上限
                elif strategy == "smart":
                    ratio = float(ai.get("sample_ratio", 0.3) or 0.3)
                    # 用文件路径哈希做稳定抽样（同一文件结果一致）
                    h = int(hashlib.md5(path.encode()).hexdigest()[:8], 16) / 0xFFFFFFFF
                    if h >= ratio:
                        skip_ai = True
                # ---- 纯色 / 低信息量图片跳过 ----
                if ext_l in IMAGE_EXT and not skip_ai:
                    try:
                        with _PILImage.open(path) as _chk:
                            _st = _ImageStat.Stat(_chk.convert("L").resize((64, 64)))
                            if _st.stddev and _st.stddev[0] < 8.0:
                                skip_ai = True
                                print(f"  [跳过AI] 纯色/低信息图: {os.path.basename(path)}")
                    except Exception:
                        pass
                # 图片＝直接送这一张；视频＝先把关键帧拼成九宫格，一次调用看完整段
                _pic, _prompt = None, ""
                if ext_l in IMAGE_EXT and not skip_ai:
                    with _PILImage.open(path) as _im:
                        _pic = _im.convert("RGB")
                    _prompt = ai.get("prompt", "")
                elif ext_l in VIDEO_EXT and not skip_ai:
                    _pic = _video_contact_sheet(path)
                    _prompt = ai.get("prompt_video") or VIDEO_SHEET_PROMPT
                if _pic is not None:
                    desc = _ai.describe_image(
                        ai["base_url"], ai.get("api_key", ""), ai["model"],
                        _pic, prompt=_prompt,
                        max_tags=int(ai.get("max_tags", 12) or 12))
                    ai_meta = {"ai_description": desc.get("description", ""),
                               "ai_tags": desc.get("tags", []),
                               "ai_model": ai["model"]}
                    n_ai += 1
                    report(phase="encoding", total=total, current=fi - 1,
                           indexed=n_new, skipped=n_skip, units=n_unit,
                           message=f"AI 描述 {os.path.basename(path)}")
            except Exception as e:
                print(f"  [AI 描述失败] {path}: {e}")

        # ASR 音频转写（仅音频，且开启配置时）
        asr_meta = {}
        if asr:
            try:
                import asr_desc as _asr
                ext_l = os.path.splitext(path)[1].lower()
                if ext_l in (AUDIO_EXT | VIDEO_EXT):     # 视频也转写（用户要的哼唱范围含视频）
                    # 时长限制（可选）
                    max_dur = asr.get("max_duration") or 0
                    dur = _probe_duration(path) if max_dur else 0
                    if max_dur and dur > 0 and dur > float(max_dur):
                        print(f"  [跳过ASR] 媒体过长 {dur:.0f}s/"
                              f"{dur/60:.1f}分钟（上限 {max_dur}s）: "
                              f"{os.path.basename(path)}")
                        asr_meta = {"asr_skip": "too_long",
                                    "asr_skip_dur": round(dur, 1)}
                    else:
                        r = _asr.transcribe(
                            asr["base_url"], asr.get("api_key", ""), asr["model"],
                            path, language=asr.get("language", ""))
                        txt = (r.get("text") or "").strip()
                        if txt:
                            asr_meta = {"asr_text": txt,
                                        "asr_model": asr["model"]}
                            n_asr += 1
                            report(phase="encoding", total=total, current=fi - 1,
                                   indexed=n_new, skipped=n_skip, units=n_unit,
                                   message=f"转写 {os.path.basename(path)}")
                        else:
                            print(f"  [ASR无内容] {os.path.basename(path)}")
                            asr_meta = {"asr_none": True}     # 确认过没内容，不再反复重试
            except Exception as e:
                # 故意不打标记：下次刷新会重试（例如 oMLX 刚启动时先失败后成功）
                print(f"  [ASR转写失败] {path}: {e}")

        for kind, idx, meta, text, image in units:
            try:
                if image is not None:
                    vec = we.embed(model, processor, image=image,
                                   instruction=we.QUERY_INSTRUCTION)
                else:
                    # 文档/音频/视频：把 AI 标签、ASR 转写文本也拼进文本一起编码
                    extra = ""
                    if ai_meta:
                        extra = " " + " ".join(ai_meta.get("ai_tags", []) or [])
                    if asr_meta and kind in ("audio", "video"):
                        extra += " " + asr_meta.get("asr_text", "")
                    vec = we.embed(model, processor, text=(text or "") + extra,
                                   instruction=we.QUERY_INSTRUCTION)
                mx.eval(vec)
                lst = vec.tolist()
                merged = dict(prev_meta)      # 底稿：保留上一轮的 asr_text / ai_tags
                merged.update(meta or {})
                merged.update(ai_meta)
                merged.update(asr_meta)
                # ★ 这次真转出文字了 → 清掉上一轮的「跳过 / 没内容」痕迹。
                #   不清的话，老记录里的 asr_skip 会一直挂在 meta 上，把这条
                #   永远挡在待转写清单外（`asr_status` 见到 asr_skip 就 continue）。
                if (asr_meta.get("asr_text") or "").strip():
                    for _k in ("asr_skip", "asr_skip_dur", "asr_none"):
                        merged.pop(_k, None)
                # 文件这次查出「没有标签」时，别把上一轮残留的 tags 留着
                if (meta or {}).get("tags_checked") and not (meta or {}).get("tags"):
                    merged.pop("tags", None)
                cur.execute(
                    "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                    " VALUES (?,?,?,?,?,?,?,?)",
                    (path, kind, idx, json.dumps(merged, ensure_ascii=False), len(lst),
                     _to_blob(lst), mtime, time.time()))
                n_unit += 1

                # ★ 每个 AI 标签各存一个「居中向量」：把标签串整体编码成一个向量时，
                #   长文本会被稀释（实测查「气球」只有 0.388，跟无关标签的 0.339 差不多），
                #   而单个标签单独编码能到 0.91。所以一个标签一条记录。
                # 视频的标签向量只写一份（在 idx==0 那一次循环里写），
                # 否则每帧都会 INSERT OR REPLACE 同一批 chunk_idx，白白重复编码。
                if image is not None and ai_meta and (kind != "video" or idx == 0):
                    _tags = [t for t in (ai_meta.get("ai_tags") or []) if t]
                    _desc = (ai_meta.get("ai_description") or "").strip()
                    _slots = [(AI_TAG_BASE + i, t, True) for i, t in enumerate(_tags)]
                    if _desc:
                        _slots.append((AI_DESC_IDX, _desc, False))
                    for _si, _stext, _is_tag in _slots:
                        try:
                            cv = _center_text_vec(model, processor, _stext)
                            if not cv:
                                continue
                            tmeta = dict(merged)
                            tmeta["_ai_tag_slot" if _is_tag else "_ai_desc_slot"] = True
                            cur.execute(
                                "INSERT OR REPLACE INTO items (path,kind,chunk_idx,meta,dim,vec,mtime,indexed_at)"
                                " VALUES (?,?,?,?,?,?,?,?)",
                                (path, kind, _si,
                                 json.dumps(tmeta, ensure_ascii=False), len(cv),
                                 _to_blob(cv), mtime, time.time()))
                            n_unit += 1
                        except Exception as e:
                            print(f"  [AI标签编码失败] {path}: {e}")
            except Exception as e:
                print(f"  [编码失败] {path} #{idx}: {e}")

        # 音频/视频：建立 Chromaprint 指纹（用于「以音搜素材」）
        try:
            ext_fp = os.path.splitext(path)[1].lower()
            if ext_fp in (AUDIO_EXT | VIDEO_EXT):
                import fingerprint as _fp
                _fp.index_file(path, "video" if ext_fp in VIDEO_EXT else "audio", mtime)
        except Exception as e:
            print(f"  [指纹索引失败] {path}: {e}")

        if units:
            if not meta_need:
                n_new += 1
        else:
            # 解不出任何单元：别谎报成「新索引」——它其实一个单元都没进库，
            # 下次刷新还会被当成新文件重试（报告里那个「新索引 1 个文件 /
            # 0 个向量」就是这么来的）。单独计数、单独提示。
            n_fail += 1
            print(f"  [无法索引] {path}")
        con.commit()
        report(phase="encoding", total=total, current=fi,
               indexed=n_new, skipped=n_skip, units=n_unit,
               message=os.path.basename(path))
        if fi % 10 == 0:
            print(f"  进度 {fi}/{total}  已索引文件 {n_new}  单元 {n_unit}")

    # ---- 清理：源里已不存在的文件 + 垃圾文件 → 从索引中移除 ----
    n_removed = 0
    try:
        n_removed = purge_stale(cur, con, roots)
        if n_removed:
            print(f"\n清理：移除 {n_removed} 个已失效/垃圾的索引记录")
    except Exception as e:
        print(f"  [清理失败] {e}")

    con.commit()
    con.close()
    dt = time.time() - t0
    rm_note = f"，清理 {n_removed} 个失效记录" if n_removed else ""
    fail_note = f"，{n_fail} 个文件无法解析" if n_fail else ""
    meta_note = f"，补元数据 {n_meta} 个" if n_meta else ""
    print(f"\n完成：新索引 {n_new} 个文件 / {n_unit} 个向量，跳过 {n_skip} 个未变文件{meta_note}{rm_note}{fail_note}，耗时 {dt:.1f}s")
    ai_note = f"，AI 描述 {n_ai} 张" if n_ai else ""
    rm_msg = f"，清理 {n_removed} 项" if n_removed else ""
    fail_msg = f"，{n_fail} 个文件无法解析" if n_fail else ""
    meta_msg = f"，补元数据 {n_meta} 个" if n_meta else ""
    report(phase="done", total=total, current=total, indexed=n_new,
           skipped=n_skip, units=n_unit, ai=n_ai, removed=n_removed, failed=n_fail,
           meta=n_meta,
           elapsed=round(dt, 1),
           message=f"新索引 {n_new} 个文件 / {n_unit} 个向量{meta_msg}{ai_note}{rm_msg}{fail_msg}")
    return {"indexed": n_new, "skipped": n_skip, "units": n_unit, "ai": n_ai,
            "asr": n_asr,
            "removed": n_removed,
            "failed": n_fail,
            "meta": n_meta,
            "total": total, "elapsed": round(dt, 1)}


# --------------------------------------------------------------------------
# 搜索
# --------------------------------------------------------------------------
def _lexical_boost(query: str, path: str, meta: dict, use_name: bool = True) -> float:
    """关键词/文件名匹配加成（模糊检索的关键：短查询靠语义分数区分度不足）。

    对文件名和路径做子串匹配，命中越多加成越高；同时支持单字级匹配，
    让「狗」这种单字查询也能把「小狗在操场奔跑.png」明显推上去。

    use_name=False 时**只保留转写文本(内容)那一档**，不做任何文件名/路径匹配。
    Agent 走这条路：文件名可能是随手起的、甚至写错的，而转写文本是真实内容。
    """
    q = query.strip().lower()
    if not q:
        return 0.0
    name = os.path.basename(path).lower()
    full = path.lower()
    # 去掉常见装饰词，提取有效字符
    stop = set("的一只在是了和与及于把被为地得着过有上下了里外中")
    chars = [c for c in q if c not in stop and not c.isspace()]
    boost = 0.0
    if use_name:
        # 1) 整串命中文件名
        if q in name:
            boost += 0.30
        elif q in full:
            boost += 0.16
        # 2) 单字/短词命中（模糊查询核心）
        if chars:
            hit = sum(1 for c in chars if c in name)
            ratio = hit / len(chars)
            boost += 0.20 * ratio
            if any(c in full for c in chars):
                boost += 0.04
    # 3) ASR 转写文本（歌词/台词）命中：整句命中给足加成，长片段命中给一半。
    #    没有这条时，搜原文歌词会输给毫不相干的空文档（WeMM 对短查询 vs 长文本
    #    区分度不足，无关文本基线就有 0.45），唱歌检索基本废掉。
    #    ★ 这一档是**内容**证据，不是文件名，所以 use_name=False 时照样保留。
    txt = ((meta.get("asr_text") or "") + " " + (meta.get("asr_clean") or "")).lower()
    if txt and len(q) >= 2:
        if q in txt:
            boost += 0.34
        else:
            try:
                import fastsearch as _fs
                run = _fs.content_run(q, txt)
            except Exception:
                run = 0
            if run:
                boost += 0.22
    return boost


def _normalize_scores(raw):
    """把原始余弦分映射到 0-1，便于阈值过滤与展示。"""
    if not raw:
        return []
    vals = [s for s, *_ in raw]
    lo, hi = min(vals), max(vals)
    if hi - lo < 1e-6:
        return [(0.5, *rest) for _, *rest in raw]
    return [((s - lo) / (hi - lo), *rest) for s, *rest in raw]


# 上位词 → 具体词 的查询扩展表（提升泛化召回）
QUERY_EXPANSION = {
    "动物": ["狗", "猫", "宠物", "鸟", "小动物"],
    "宠物": ["狗", "猫", "小动物", "宠物"],
    "犬": ["狗", "小狗", "犬"],
    "狗": ["狗", "小狗", "狗狗", "宠物"],
    "猫": ["猫", "小猫", "猫咪", "宠物"],
    "人": ["人", "人物", "男子", "女子", "人像"],
    "人物": ["人", "人物", "男子", "女子", "人像"],
    "食物": ["食物", "美食", "菜", "餐"],
    "美食": ["美食", "食物", "菜", "餐"],
    "风景": ["风景", "景色", "自然", "山", "天空"],
    "景色": ["风景", "景色", "自然", "天空"],
    "自然": ["自然", "风景", "山", "树", "天空"],
    "车": ["车", "汽车", "车辆", "轿车"],
    "建筑": ["建筑", "房子", "楼", "房屋"],
    "植物": ["植物", "花", "树", "草"],
    "文字": ["文字", "文本", "文档", "报告"],
    "文档": ["文档", "文本", "报告", "文件"],
    "报表": ["报表", "报告", "财务", "数据"],
}


def expand_query(query: str) -> list:
    """把查询扩展成多路检索词（含原词）。上位词补齐具体实例词以提升泛化召回。"""
    q = (query or "").strip()
    if not q:
        return []
    variants = [q]
    if q in QUERY_EXPANSION:
        variants.extend(QUERY_EXPANSION[q])
    else:
        for key, exts in QUERY_EXPANSION.items():
            if key in q and key != q:
                variants.extend(exts); break
    seen, out = set(), []
    for v in variants:
        if v and v not in seen:
            seen.add(v); out.append(v)
    return out


def _is_meaningless(path: str, meta: dict) -> bool:
    """判断是否为「无意义」素材（纯色/测试/占位图），用于降权。

    依据：文件名关键词、文件体积过小、图片色彩单一（纯色图）。
    """
    try:
        name = os.path.basename(path).lower()
        for kw in ("testsrc", "test_pattern", "solid", "blank", "color=", "redshot"):
            if kw in name:
                return True
        size = meta.get("size") or 0
        if not size:
            try: size = os.path.getsize(path)
            except OSError: size = 0
        if size and size < 8192:
            return True
        # 图片：检测是否近似纯色（标准差极小）
        if os.path.splitext(path)[1].lower() in IMAGE_EXT:
            try:
                from PIL import Image, ImageStat
                with Image.open(path) as im:
                    st = ImageStat.Stat(im.convert("L").resize((64, 64)))
                    if st.stddev and st.stddev[0] < 6.0:
                        return True      # 近乎纯色
            except Exception:
                pass
    except Exception:
        pass
    return False


def _low_info_doc(meta: dict) -> bool:
    """判断是否为「没有正文」的文档（扫描件、老式二进制、空模板、几个字的占位文件）。

    这类文档的文本向量等于「一个空文档」，落在向量空间中心附近，与**任何**
    查询都有 0.5~0.64 的余弦相似度——语义排序里它们永远是前排，且与查询内容
    无关。实测：搜「太空」「美女」「小女生」，前 5 名都是同一批空文档。

    `text_len` / `no_text` 由 iter_units 在索引时写进 meta；老索引没有这两个
    字段（返回 None）时不做判断，避免误伤。
    """
    if not meta or meta.get("text_len") is None:
        return False
    return bool(meta.get("no_text")) or meta["text_len"] < EMPTY_TEXT_LEN


def _demote_low_info(results):
    """把「没有正文」的文档从**纯向量**排序里压下去。

    results: [(score, path, kind, chunk_idx, meta), ...]

    空文档的向量对任何查询都偏高——实测搜「唱歌」时 `未命名.html`（去掉 HTML 标签后
    正文只剩 5 个字符）拿到 0.754，比真正的答案 `告白_FR.mp3`（0.503）还高。这里只压
    向量分；L0 文件名命中的分数是在合并阶段另外取的，不受影响（搜「6」仍能找到 6.wps）。
    """
    out = [(s * EMPTY_DOC_PENALTY if (k == "document" and _low_info_doc(m)) else s, p, k, i, m)
           for s, p, k, i, m in results]
    out.sort(key=lambda x: -x[0])
    return out


def search(query, top_k=10, kind=None, db_path=DB_PATH, model_path=None,
           threshold=0.0, fuzzy=True, tier="auto", lexical=True):
    """三级递进语义检索（够用即止，省资源）。

    lexical=False 时不做文件名/路径加成，只按内容相似度算分（Agent 用这条路，
    理由见 `_lexical_boost` 的 docstring）。

    tier:
      "fast"   —— 只用 L0 文件名匹配（0 模型调用）
      "vector" —— L0 + L1 向量检索（用缓存的查询向量，不加载模型）
      "deep"   —— 强制走 WeMM 深度语义（查询扩展 + 多指令融合）
      "auto"   —— 自动：L0 → L1（缓存命中）→ 必要时 L2 深度

    返回 [(score, path, kind, chunk_idx, meta)]
    """
    import fastsearch as fs

    # ---------- L0：文件名精确匹配（零模型调用） ----------
    if tier in ("auto", "fast"):
        lex = fs.search_lexical(query, db_path, kind=kind)
        if tier == "fast":
            return [(s, p, k, i, m) for s, p, k, i, m in lex[:top_k]]
        # auto：命中且质量足够 → 直接返回（简单查询不碰模型）
        if lex and fs.score_quality(lex, min_count=1, min_top=0.80) and len(lex) >= 1:
            if lex[0][0] >= 0.85:
                return [(s, p, k, i, m) for s, p, k, i, m in lex[:top_k]]

    # ---------- L1：向量检索（优先用缓存，不加载模型） ----------
    if tier in ("auto", "vector"):
        qv = fs.qcache_get(query, "") or fs.qcache_get(query, "q")
        if qv is None and tier == "vector":
            # vector 模式：需要编码查询，但用缓存避免重复
            qv = _encode_query_cached(query)
        if qv is not None:
            # with_chunk=True 才会回填 meta —— _demote_low_info 要靠 meta 里的
            # text_len/no_text 判断「空文档」，也顺带让视频定位到最像的那一帧。
            vec_res = _demote_low_info(
                fs.search_vector(qv, db_path, kind=kind, with_chunk=True))
            if vec_res:
                # 与 L0 结果融合
                merged = {}
                for s, p, k, i, m in lex if tier == "auto" else []:
                    merged[p] = (s, p, k, i, m)
                for s, p, k, i, m in vec_res:
                    if p in merged:
                        merged[p] = (max(merged[p][0], s) + 0.1, p, k, i, m)
                    else:
                        merged[p] = (s, p, k, i, m)
                res = sorted(merged.values(), key=lambda x: -x[0])
                if tier == "vector":
                    return res[:top_k]
                # L1「够好」的判定必须跟界面门槛对齐（ui.html 的 MIN_RELEVANT = 0.46）：
                # 如果 L1 的最高分连门槛都够不着，交上去也只会被前端全部隐藏成
                # 「没有找到相关素材」—— 那种情况应该继续下沉到 L2 深度语义。
                if fs.score_quality(res, min_count=2, min_top=0.46):
                    return res[:top_k]

    # ---------- L2：WeMM 深度语义 ----------
    import mlx.core as mx
    import embed as we

    model_path = model_path or we.DEFAULT_MODEL
    model, processor = we.load_model(model_path)   # 带全局缓存，不会重复加载

    terms = expand_query(query) if fuzzy else [query]
    instrs = [we.QUERY_INSTRUCTION, ""] if fuzzy else [we.QUERY_INSTRUCTION]

    qvecs, qvecs_extra = [], []
    for i, term in enumerate(terms):
        for instr in (instrs if i == 0 else [we.QUERY_INSTRUCTION]):
            v = we.embed(model, processor, text=term, instruction=instr)
            mx.eval(v)
            vec = v.tolist()
            (qvecs if i == 0 else qvecs_extra).append(vec)
            # 缓存查询向量（供 L1 复用）
            if i == 0 and instr == instrs[0]:
                fs.qcache_put(term, "", vec)

    con = sqlite3.connect(db_path)
    _kf, _ka = fs._kind_sql(kind)
    sql = ("SELECT path,kind,chunk_idx,meta,dim,vec FROM items WHERE 1=1" + _kf)
    try:
        rows = con.execute(sql, _ka).fetchall()
    except sqlite3.OperationalError:
        rows = []
    finally:
        con.close()

    agg = {}
    for path, k, idx, meta, dim, blob in rows:
        v = _from_blob(blob)
        if len(v) != len(qvecs[0]):
            continue
        sim_main = max(sum(a * b for a, b in zip(q, v)) for q in qvecs)
        sim_ext = 0.0
        if qvecs_extra:
            sim_ext = max(sum(a * b for a, b in zip(q, v)) for q in qvecs_extra)
        # 帧定位用「主查询向量」的相似度。L1 只能拿到缓存的那一份（QUERY_INSTRUCTION），
        # 若这里改用 sim_main（含无指令那一份、扩展词的最大值），同一个查询在
        # tier=vector 和 tier=deep 下会定位到不同的帧 —— 两处口径必须一致。
        sim_frame = sum(a * b for a, b in zip(qvecs[0], v))
        # AI 标签/描述向量（chunk_idx >= 9000）：居中点积 → 映射成与文档可比的分数。
        # 只用主指令（QUERY_INSTRUCTION）那一份 —— 映射常数是按它标定的，
        # 无指令那一份的尺度不同，混进来会把无关查询抬过门槛。
        if idx >= AI_TAG_BASE:
            sim_main = fs.map_tag_score(sim_main)
            sim_ext = 0.0
        m = json.loads(meta or "{}")
        # 定位用的「最匹配帧」：只在真实帧行（chunk_idx < AI_TAG_BASE）里选。
        # 视频的 AI 标签向量是整段拼图生成的，不对应任何一帧，用它当命中帧会把封面钉在片头。
        is_frame = idx < AI_TAG_BASE
        if path not in agg:
            agg[path] = {"path": path, "kind": k, "idx": idx, "meta": m,
                         "sim_main": sim_main, "sim_ext": sim_ext,
                         "fidx": idx if is_frame else None,
                         "fmeta": m if is_frame else None,
                         "fsim": sim_frame if is_frame else None}
        else:
            a = agg[path]
            if sim_main > a["sim_main"]:
                a["sim_main"], a["idx"], a["meta"] = sim_main, idx, m
            if is_frame and (a["fsim"] is None or sim_frame > a["fsim"]):
                a["fsim"], a["fidx"], a["fmeta"] = sim_frame, idx, m
            a["sim_ext"] = max(a["sim_ext"], sim_ext)

    scored = []
    for a in agg.values():
        sm, se = a["sim_main"], a["sim_ext"]
        sim = max(sm, se * 0.98)
        lex_b = _lexical_boost(query, a["path"], a["meta"], use_name=lexical) if fuzzy else 0.0
        total = sim + lex_b
        if fuzzy and _is_meaningless(a["path"], a["meta"]):
            total *= 0.55
        if fuzzy and _low_info_doc(a["meta"]):
            # 空文档只罚语义分：它的 sim 是「通用基线」不是「相关」；
            # 但文件名命中的 lex_b 是真实证据（搜「6」仍要能找到 6.wps），保留。
            total = sim * EMPTY_DOC_PENALTY + lex_b
        # 对外报的是「命中帧」的 chunk_idx 与 meta（含 timestamp），分数仍是全局最高分
        oidx = a["fidx"] if a["fidx"] is not None else a["idx"]
        ometa = a["fmeta"] if a["fidx"] is not None else a["meta"]
        scored.append((total, sim, a["path"], a["kind"], oidx, ometa))

    scored.sort(key=lambda x: -x[0])
    if fuzzy and scored:
        hi = scored[0][0]
        cutoff = hi * 0.42 if hi > 0 else -1e9
        if threshold > 0:
            cutoff = max(cutoff, threshold)
        scored = [x for x in scored if x[0] >= cutoff]

    return [(total, path, k, idx, m) for total, sim, path, k, idx, m in scored[:top_k]]


def _encode_query_cached(query: str):
    """编码查询向量并写入 L1 缓存（不加载模型时返回 None）。"""
    import fastsearch as fs
    hit = fs.qcache_get(query, "")
    if hit is not None:
        return hit
    try:
        import mlx.core as mx
        import embed as we
        model, processor = we.load_model(we.DEFAULT_MODEL)
        v = we.embed(model, processor, text=query, instruction=we.QUERY_INSTRUCTION)
        mx.eval(v)
        vec = v.tolist()
        fs.qcache_put(query, "", vec)
        return vec
    except Exception:
        return None


def stats(db_path=DB_PATH):
    con = sqlite3.connect(db_path)
    try:
        total = con.execute("SELECT COUNT(*) FROM items").fetchone()[0]
        files = con.execute("SELECT COUNT(DISTINCT path) FROM items").fetchone()[0]
        by_kind = con.execute("SELECT kind, COUNT(*) FROM items GROUP BY kind").fetchall()
    except sqlite3.OperationalError:
        total, files, by_kind = 0, 0, []
    finally:
        con.close()
    return {"vectors": total, "files": files, "by_kind": dict(by_kind)}


def main():
    ap = argparse.ArgumentParser(description="本地语义检索（图片/视频/文档）")
    sub = ap.add_subparsers(dest="cmd", required=True)

    p_idx = sub.add_parser("index", help="建立/更新索引")
    p_idx.add_argument("roots", nargs="+", help="目录或文件")
    p_idx.add_argument("--force", action="store_true", help="强制重建")
    p_idx.add_argument("--model", default=None)

    p_s = sub.add_parser("search", help="搜索")
    p_s.add_argument("query")
    p_s.add_argument("-k", "--top-k", type=int, default=10)
    p_s.add_argument("--kind", choices=["image", "video", "document"])
    p_s.add_argument("--model", default=None)

    sub.add_parser("stats", help="索引统计")

    args = ap.parse_args()
    if args.cmd == "index":
        build_index(args.roots, force=args.force, model_path=args.model)
    elif args.cmd == "search":
        res = search(args.query, top_k=args.top_k, kind=args.kind, model_path=args.model)
        if not res:
            print("(无结果)")
        for s, path, k, idx, meta in res:
            tag = f"{k}#{idx}" if k != "image" else "image"
            extra = f" t={meta['timestamp']}s" if "timestamp" in meta else ""
            print(f"{s:.4f}  [{tag}]{extra}  {path}")
    elif args.cmd == "stats":
        print(json.dumps(stats(), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
